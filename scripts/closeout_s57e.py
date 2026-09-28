"""Session 57's evening batch, into both sheets of the close-out workbook."""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

CHECKLIST = [
    ("The overnight queue and its Stop hook", "Done", "2026-09-28",
     "28 items split into steps in docs/progress.md, in her priority order. The hook blocks the end of a turn while work remains. THREE WAYS OUT, all proved: a .claude/overnight-off file, a ceiling of 60 continuations, and an item that has not moved across three stops marked blocked in the file so the run passes it. Fails open - every error path exits silently. 11 checks against a throwaway repo copy so the destructive paths could be exercised safely."),

    ("Onboarding never sets the goal", "Diagnosed", "2026-09-28",
     "Nothing in onboarding writes fat_focus_state or muscle_focus_state, and asFocus turns a missing value into 'maintain'. So EVERY user finishes onboarding on a maintenance target whatever she said she wanted; her goal is stored as a sentence in user_context and no part of the target arithmetic reads it. The machinery and the front door are one column apart."),
    ("A correction to the brief, before anything is built", "Diagnosed", "2026-09-28",
     "My Week, My Rules, My Plans by goal and the Now/Next/Goal ladder are in the SPECIFICATION and not in the code - searched by name and by every plausible identifier. The Plans tab, conversely, left the Almanac on 20 September and the spec still says otherwise. The onboarding she wants feeds two things that do not exist, so it is build work rather than wiring."),
    ("New onboarding proposed, nine screens", "Proposed - needs her", "2026-09-28",
     "About a minute of taps then one conversation, replacing eight conversational screens. Exact wording, and what every answer writes. Screen 1 is the join that does not exist: direction -> focus states -> targets, confirmed in plain words on the same screen. It ends with targets, a first plan and the rest-of-today button rather than a welcome."),
    ("Principles 3 and 9, and Part Eleven", "Proposed - needs her", "2026-09-28",
     "Graduation never meant the app makes itself unnecessary. Rewritten as: never trap, build literacy that lasts, be there for life. Principle 9 becomes reactive by default and scheduled when asked, with Guide me / Let me lead, and neither mode keeping score."),

    ("THE APP TELLS A POST-MENOPAUSAL WOMAN SHE IS ON CYCLE DAY 1826", "Diagnosed", "2026-09-28",
     "Demonstrated with scripts/probe-cycle-no-periods.mjs, not argued. The count has no upper bound, everything past day 16 is labelled luteal, and past day 22 her weight reading is explained as water retention. It only stays quiet if she has NEVER logged a period, so somebody whose periods stop gets more wrong every day. Nothing in the app knows what HRT is. For an app whose whole audience is women over 40 this is not an edge case."),
    ("Menopause research, with citations", "Done", "2026-09-28",
     "NICE NG23 (diagnose from symptoms alone over 45, no blood tests), NHS (symptoms, 7-9 years), and the BMS/WHC factsheets with figures: 1.5 kg a year through the transition, visceral fat 5-8% to 10-15%, calcium 700 or 1200 mg, a 500 kcal deficit, strength work 'almost non-negotiable'. The BMS line on respectful consent is principle 9 written by a menopause society."),
    ("Where Selodia could actually beat Balance", "Done", "2026-09-28",
     "Every menopause app tracks symptoms against TIME. None can answer 'is my sleep worse in the week my joints ache' because none of them hold food, movement and sleep in the same record. Selodia already does. Adding symptoms makes it the only app that can put a hot flush beside a bad night beside a heavy week."),

    ("The cost model was 86% wrong", "Fixed", "2026-09-28",
     "Measured 3.30c a chat turn against 1.77c estimated, because the estimate assumed 65 output tokens on the classify call and the real average is 132. All in with overheads: GBP 3.87 a user at a thousand subscribers, GBP 7.53 at twenty-five. GBP 7 would have lost money on every early subscriber. The classify call is 83% of the bill and is the cheapest lever available."),
    ("Price re-modelled", "Proposed - needs her", "2026-09-28",
     "GBP 12/month and GBP 120/year, founding GBP 8 / GBP 79 for the first 100 kept for good. Break-even fifteen subscribers at GBP 12, eleven at GBP 15. GBP 15 is where this goes after guided sessions and a coach, not before. Nikki's two numbers are the real constraint: GBP 10 for a yoga app with five videos, GBP 2 for tracking."),

    ("Premium means restraint - standing brand rule", "Done", "2026-09-28",
     "Six absolute rules in the decision patterns with the evidence, and in the spec as a checkable list. No escalating or personalised discounts; no fake urgency; email only what is needed or asked for; no 'as featured in' unless true and linked; no body images; goals are the user's words. The test: would this still be here if it did not increase conversion?"),
    ("Selodia checked against its own rule", "Done", "2026-09-28",
     "Nothing breaks it. Six checks, all clean - and the six urgency matches were all false, five signed-URL expiry and one a code comment. Two things named that are not breaches: the goal is stored in her own words and then read by nothing, and the password-reset email still does not work. Rule 1 has no code to break yet, which is exactly when it will be broken."),

    ("Meal suggestions shipped and verified", "Done", "2026-09-28",
     "The writer now receives today's targets and a fortnight of her own meals. buildDayStatePrompt had run on every turn for months and gone to the classify call only - the half that answers had never seen it. Both writer call sites carry the new fields, because adding it to one would have worked on turns that report nothing and silently not on turns that report something."),
    ("An empty day is not a missing target", "Fixed", "2026-09-28",
     "The live app said 'there's no target set to work from' on a turn where the record said 1,760 kcal and 83-99g. Asked directly for the target two minutes later, the same account answered correctly. The figures were there and the reply denied them. One clause, deployed and verified."),
    ("Three meal cases in the test set", "Done", "2026-09-28",
     "51/51 on the rebuilt prompt, 42/51 on the old one - the old path fails 'uses the figures it was given' immediately, which is the regression they exist to catch. The self-test refused them twice before a token was spent."),
    ("A check that tested a wording, not a behaviour", "Fixed", "2026-09-28",
     "'does not ask her to retype' fired on a reply saying 'No need to type it again', which is the app doing the right thing in the words the check was hunting for. A negative check hides this, because the temptation is to soften the reply until the check is happy - which trains the prompt to avoid a word rather than a behaviour."),
    ("A missing demonstration says so", "Done", "2026-09-28",
     "Her decision, reversing a documented one, and the old reasoning is kept rather than deleted. A blank where every other exercise has a clip does not read as 'we cover patterns, not names'; it reads as broken."),
    ("The spec claimed there was no video player", "Fixed", "2026-09-28",
     "Stale for weeks. There is one, it works, and it is rendered against every exercise in a saved plan. What IS missing is beginner calisthenics coverage, which is content rather than plumbing."),
]

DECISIONS = [
    ("A default that silently stands in for a missing answer will be mistaken for an answer", "2026-09-28",
     "asFocus turns a missing focus into 'maintain'. Nothing is null downstream, nothing errors, nothing logs - every user simply comes out on a maintenance target. What makes it hard to see is that the default is REASONABLE; a reviewer reading that line nods. The bug is not the value, it is that 'she chose maintain' and 'nobody ever asked her' produce identical state, so nothing afterwards can tell them apart. Keep them apart in the data even when they behave the same in the code."),
    ("A guard that can only say keep going is a trap, so build the ways out first and prove each one", "2026-09-28",
     "The stop hook blocks the end of a turn and writes to the progress file, so a bug in it could have run unattended until morning. Three escapes went in before it was enabled, and it fails open on every error path. The part worth generalising is the testing: each valve was proved against a throwaway copy of the repo so the destructive paths could be exercised safely. When a mechanism's job is to keep something running, its tests have to be about stopping."),
    ("A check that matches a phrase is testing wording; one that matches an intent is testing behaviour", "2026-09-28",
     "'does not ask her to retype' failed a reply that said 'No need to type it again'. A negative check hides this because the failure looks like a pass in reverse, and the tempting fix is to soften the reply until the check is happy - which trains the prompt to avoid a word rather than a behaviour. When a check fires, read the sentence it fired on before changing anything else. If the sentence is good, the check is wrong."),
    ("Write the brand rule down before there is any pressure to break it", "2026-09-28",
     "Six rules recorded on a day when the app has no billing code at all, which is the only moment it costs nothing to agree. Every one of them is something a competitor does profitably, and each will arrive later as a single checkbox in a store dashboard labelled 'introductory offer' or 'win-back campaign'. The test that makes it operable rather than aspirational: would this still be here if it did not increase conversion?"),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 1
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value.startswith(("Done", "Fixed", "Diagnosed")):
                    cell.fill = copy(done_fill)
            r += 1

    append(check, "S57 continued: 28Sep26  (evening batch)", CHECKLIST, 2)
    append(dec, "S57 continued: 28Sep26  (evening batch)", DECISIONS, None)
    wb.save(BOOK)
    print(f"Checklist {check.max_row} rows, Decisions {dec.max_row} rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
