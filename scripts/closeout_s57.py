"""Append Session 57 to both sheets of the close-out workbook.

Two continuous sheets, appended to. Never a tab per session.
"""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

CHECKLIST = [
    ("The chat switch is on", "Done", "2026-09-28",
     "Deployed 09:55 after she read the before-and-after replies. Confirmed live from HER OWN ROWS rather than a deploy log: every reply from 10:17 carries an answers_id, which only the new path sets."),
    ("Every reply since 20:56 on the 27th was thrown away", "Done", "2026-09-28",
     "My bug, shipped seven minutes before the first lost turn. ON CONFLICT (answers_id) cannot infer a PARTIAL index, so the insert raised 42P10 on every call. The route logs a failed write and carries on - deliberately - so she saw her replies and nothing was stored. Three lost. Plain insert now; 23505 is the expected no-op."),
    ("The index was verified and the call site was not", "Diagnosed", "2026-09-28",
     "The lesson rather than the bug. One statement run against the real schema would have shown it in ten seconds. A constraint the database accepts is not a constraint the application can use."),
    ("The new path could have rewritten a C-SSRS screening question", "Done", "2026-09-28",
     "Found hours after go-live by reading applySafetyStateMachine for something else. The switch ran on every turn the ALLERGY gate passed, and the safety machine substitutes a FIXED screening question precisely so a probe cannot co-occur with a card. turnIsOrdinary() is now a whitelist with ten checks. Not caught by any test: the chat set has a medical case and no distress case."),
    ("Fallbacks to the old path are recorded", "Done", "2026-09-28",
     "reply_path_fallbacks: reason (error/empty/max_tokens), detail, voice, turn. A console line lives in Vercel where 'how often is this happening' cannot be answered."),
    ("The 10:32 fallback, proven by contradiction", "Diagnosed", "2026-09-28",
     "That reply carried the appended 'Kept in your Almanac, under Insights, as a symptom'. The new path clears appended notes, so the line can only appear if the writer returned nothing. 3 seconds against 10 for the turns around it. WHY it failed is unknown and only the new table can answer it, for turns from now on."),

    # ── Item 13, the allergy gate ─────────────────────────────────────────────
    ("13a: the gate blocked two plain questions about nickel", "Done", "2026-09-28",
     "Layer 3 matched allergen names against the reply REGARDLESS of whether it suggested food, deliberately, because a self-report can be gamed. Right about self-reports, wrong about what a match proves: every true answer to a question about nickel contains the word. A hit without a food flag is now adjudicated by layer 4 rather than blocked."),
    ("13b: a contact allergy was a food restriction", "Done", "2026-09-28",
     "allergies.kind - food, contact, environmental, other. Only food and other arm the filter. What was in her table: 'seasonal allergy (hay fever, summer)' and 'nickel'. Neither is food; both were arming a food gate, one for eight days."),
    ("13c: filed as a symptom in Insights, not Me", "Done", "2026-09-28",
     "A proposal's type is fixed when the OFFER is made, so 'yes, perhaps under skincare and allergies' accepted a symptom decided fifteen minutes earlier. applyRedirect re-aims a waiting offer. Her entry moved to Almanac > Me, section 'Skincare and allergies'."),
    ("13d: fourteen checks, and the switch was on at 10:33", "Done", "2026-09-28",
     "It was. The blocked replies came from the new path - but a blocked reply is not the chat path's output: the gate replaces the reply AFTER it is written."),

    # ── Item 12 ───────────────────────────────────────────────────────────────
    ("12a: the weigh-in card compared against a six-day-old reading", "Done", "2026-09-28",
     "56.95 that morning, 56.85 the morning before, 55.55 six days earlier. It reported +1.4 kg over a day-on-day change of +0.1, because findWeekAgoReading finds the reading NEAREST SEVEN DAYS BACK. Right for the Overview, wrong for a card shown as she steps off the scale."),
    ("12a: 'worth a calm look rather than a shrug or a spiral'", "Done", "2026-09-28",
     "Over a hundred grams. Reverses a deliberate decision that a trend outranks the noise flags - true that rebutting mid-sentence reads as an argument with itself, and it decided which half wins. The explanation now leads and the sentence stops. Her words: 'It's nothing to be alarmed about.'"),
    ("A big day is a flag at all, for the first time", "Done", "2026-09-28",
     "There was a sodium flag and nothing for having simply eaten a lot, which is the most ordinary reason a scale is up. High bar, short window: a 1,400 kcal Tuesday does not trip it."),
    ("12b: three copies of Sunday's pizza", "Done", "2026-09-28",
     "The real one at 21:04 Sunday, a copy logged for Monday, and a third from the date correction which moved it to Sunday MORNING and removed neither. Backed up to a new food_logs_removed table and deleted. She reported it again later; the database was already clean, and the Log tab does refetch on focus - the likeliest explanation is a screen she had not left and returned to."),
    ("12d: half a lager at 6 g protein", "Done", "2026-09-28",
     "CoFID's 'Lager, standard' is 24 kcal and 0.3 g per 100 ml, so a 284 ml half is 68 kcal and 0.9 g. Calories 2.6x out, protein 7x. A 10 ml splash of whole milk was logged at 2.5 g protein. Her entries corrected; the estimator is not."),

    # ── Items 1 to 11 ─────────────────────────────────────────────────────────
    ("Build log: tracked time and locations for S45-56", "Done", "2026-09-28",
     "Her timer totals; locations for S50-S56; start times left unrecorded at her instruction. S54 and S56 corrected to Wood Street Library and the V&A East Storehouse."),
    ("Open Actionables reconciled against the live app", "Done", "2026-09-28",
     "Three items listed as not started had shipped the same day the list was written, in my own commits. A fourth turned up an hour later. The cause: the list was rebuilt from the LIST rather than from the app. Every line now records how it was checked."),
    ("Botanicals", "Done", "2026-09-28",
     "BotanicalMark was already deleted in b4e2090; the Sprig on Settings stays. Nothing to do."),
    ("The 09:24 duplicate reply, and the gap", "Done", "2026-09-28",
     "Backed up and deleted; zero repeated assistant rows inside fifteen seconds across 1,012. The turn id is now minted in code, so a chat reply always has something to answer and the partial index always applies."),
    ("6: the roundup's content half", "Done", "2026-09-28",
     "Rebuilt from a baseline in app/lib/roundup-prompt.ts. The numbered ORDER list and its 'one thematic observation' step are gone - that instruction is where 'the thread running through this week is permission' came from, and it produced one every time on the same week. Tested on HER real week: 15/15."),
    ("A scale read six times in 28 seconds is one weigh-in", "Done", "2026-09-28",
     "Six of her eight readings that week were a cluster on the 24th - 58.0 twice then 55.5 four times. Counted raw that is eight readings and a delta across values 2.5 kg apart. Collapsed at the read; nothing deleted."),
    ("7: an ordinary food log got a reply that could follow anything", "Done", "2026-09-28",
     "Saying nothing is not the alternative to a receipt - both prove nobody was listening. Chat test set 35/35 on the new path, 29/35 on the old."),
    ("8: a saved plan logged as an activity", "Done", "2026-09-28",
     "The session writer exists and was never reached - the thigh failure in a different tab. Guard at the write, stricter than the deliberate path because nothing there has decided a routine happened. 13 checks."),
    ("9: repeated-phrase probe within one conversation", "Done", "2026-09-28",
     "ALREADY BUILT - the fourth item today found done while listed as not started. Now splits at the rebuild: 'got it' was 27% of 359 replies before, and the commonest opening since is 10% of 42."),
    ("10: voice speed", "Done", "2026-09-28",
     "Measured from her rows: 3.3s median on the 26th (49 turns) and the 27th (32 turns), 7.0s today (8 turns). So yes, turns WERE back to 3.2s. The new path doubled it. The classify call was still writing a full reply nobody reads; it is asked for short now, except on a distress turn where it is still the real reply."),
    ("11: research, no building", "Done", "2026-09-28",
     "Four documents in Build Specs dated 2026-09-28: store rules, ICO registration, the beta agreement first draft, and billing plus complimentary access plus wave one. Every one says where I could not establish something."),
    ("A feedback route that carries context", "Done", "2026-09-28",
     "The last unblocked item on the Open Actionables list and one of wave one's three blockers. Settings > 'Something not right?', writing to feedback_reports with the update id, platform and OS attached. What is attached is listed, not summarised."),
    ("Today's four migrations written into the repo", "Done", "2026-09-28",
     "food_logs_removed, allergies.kind, reply_path_fallbacks, feedback_reports. All applied through the management API and existing only in the database until now."),

    # ── Hers, and carried ─────────────────────────────────────────────────────
    ("Apple Developer enrolment", "Outstanding - with Apple", "2026-09-28",
     "Submitted for Selodia Ltd, enrolment ID 3N9H5LB49A, using the D-U-N-S from 10 September. Sign in with Apple and an ios.bundleIdentifier follow the account. Neither beta wave is affected: both are Android."),
    ("Supabase auth config: redirect URL and email template", "Outstanding - needs her", "2026-09-28",
     "Carried. Needs a personal access token (sbp_), not the service role key. Required before wave one."),
    ("The turn latency", "Outstanding", "2026-09-28",
     "3.3s to 7.0s. Partly recovered; the rest is inherent to a second call. Needs re-measuring on real turns tomorrow."),
    ("Six junk weight rows on 24 September", "Outstanding - needs her", "2026-09-28",
     "58.0 twice and 55.5 four times within 28 seconds. The roundup now collapses them, so nothing reads wrong - but deleting body measurements needs her say-so."),
    ("Drinks estimator", "Outstanding", "2026-09-28",
     "Her entries are corrected and the estimator is not. CoFID is already in the database with the real figures, which makes a drinks-only lookup a narrow and well-justified case."),
    ("Removing the old chat path", "Queued", "2026-09-28",
     "Still there behind the switch, which is right for now. Beta cannot start with two paths."),
    ("Session close-out ceremony", "Done", "2026-09-28",
     "Workbook, build log, article log, decision patterns, spec, beta checklist and Open Actionables all written; five documents synced and verified; tree clean and pushed."),
]

DECISIONS = [
    ("A constraint the database accepts is not a constraint the application can use", "2026-09-28",
     "A partial unique index was applied, verified and correct. The client call asked Postgres to ignore conflicts on that column, and a partial index cannot be inferred from an ON CONFLICT with no matching WHERE - so it errored on EVERY call rather than only on a conflict. The route logs a failed write and carries on, deliberately, and that deliberate kindness hid it for thirteen hours. The mistake was verifying the migration and not the call site: checking that a thing exists rather than that the code can use it."),
    ("A check on what was SAID is not a check on what was MEANT", "2026-09-28",
     "The allergy gate matched allergen names against reply text on purpose, ignoring the model's self-report, because a string comparison cannot be gamed. Right about self-reports and wrong about what a match proves: it read 'mentions nickel' as 'suggests nickel', and those come apart on exactly one subject - questions about the allergy itself. A deterministic backstop still has to be a backstop against the right thing."),
    ("Ask what the question is before reusing the function that answers a different one", "2026-09-28",
     "The weigh-in card shared the Overview's week-ago reference so two surfaces could not caption the same comparison differently. They are not the same comparison: one asks how this week compares with last, the other asks what changed since last time. It reported +1.4 kg over a day-on-day +0.1. And when the record explains a rise, the explanation should lead and the sentence should stop - naming two ways of overreacting is still a sentence about how to feel."),
    ("A status document drifts because checking it is a different job from updating it", "2026-09-28",
     "Four items on the open list were already built, three of them shipped the same day the list was written, in the same person's commits. No carelessness: the list was rebuilt from the previous list, faithfully, and every stale line was carried forward intact. The tell is an update that never removes anything. A list that only grows is one nobody is checking against reality, because reality removes things."),
    ("Measure the cost of the fix as well as the benefit, on real traffic", "2026-09-28",
     "The new chat path doubled the median turn from 3.3s to 7.0s, measured from stored rows - offline timings were useless, since a call with no system prompt at all measured 3.4s in the same run. What the real measurement found was the waste: the first call still generating a full reply the new path discards. A feature that is better and slower is a trade. Better and slower partly for no reason is unfinished, and only a number taken after shipping shows the difference."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 1
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value in ("Done", "Diagnosed", "Fixed"):
                    cell.fill = copy(done_fill)
            r += 1

    append(check, "S57: 28Sep26  (09:40 start, Wood Street Library, mobile only)", CHECKLIST, 2)
    append(dec, "S57: 28Sep26", DECISIONS, None)
    wb.save(BOOK)
    print(f"Checklist {check.max_row} rows, Decisions {dec.max_row} rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
