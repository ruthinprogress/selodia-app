"""She edited the email before sending: it went out as Ruth Christianson-Monroy.

I had left the signature as "Ruth Christianson" and written a note saying that
was what the recipient received. The REASONING was right - a record of a sent
email must describe what actually went out - but the fact was wrong, and the
same reasoning now says the opposite: the document should be corrected, because
the version on file no longer matches the email.

Three places carry the wrong claim, and all three are fixed here:
  the email document's signature and its sent note,
  the Open Actionables row,
  the close-out workbook row.
"""

import shutil
import sys
from pathlib import Path

import docx
import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"H:\My Drive\Selodia App Project Master Folder")
FEMTECH = ROOT / "Build Specs" / "2026-09-28 Accelerating FemTech email - for approval.docx"
ACTIONS = ROOT / "Build Specs" / "Selodia - Open Actionables.docx"
BOOK = ROOT / "Build Specs" / "Selodia-Session-Closeouts.xlsx"

BARE = "Ruth Christianson"
FULL = "Ruth Christianson-Monroy"

SENT_NOTE = (
    "SENT on Tuesday 29 September 2026, from hello@selodia.app to "
    "hin.southlondon@nhs.net. Ruth edited the signature before sending, so it "
    "went out as Ruth Christianson-Monroy, and the draft below has been "
    "corrected to match what was actually sent."
)

DONE_NOTE = (
    "Done. Sent from hello@selodia.app to hin.southlondon@nhs.net, asking when the early-stage "
    "programme reopens and whether there is a list to be on. hello@selodia.app is set up as a Send As "
    "address in Gmail and both sends and receives, so the earlier caution about the from line is "
    "resolved and should not be raised again. She edited the signature before sending: it went out as "
    "Ruth Christianson-Monroy. NEXT: a reply, or nothing. If nothing by the end of October, chase once."
)


def set_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for run in p.runs[1:]:
        run._element.getparent().remove(run._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def replace_in_paragraph(par, old, new) -> int:
    if old not in par.text or new in par.text:
        return 0
    for run in par.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return 1
    text = par.text.replace(old, new)
    for run in par.runs[1:]:
        run._element.getparent().remove(run._element)
    if par.runs:
        par.runs[0].text = text
    return 1


def fix_email() -> bool:
    shutil.copy2(FEMTECH, FEMTECH.with_name(FEMTECH.stem + " (backup 2026-09-29b).docx"))
    d = docx.Document(str(FEMTECH))
    n = 0
    for par in d.paragraphs:
        if par.text.strip().startswith("SENT on"):
            for run in par.runs[1:]:
                run._element.getparent().remove(run._element)
            if par.runs:
                par.runs[0].text = SENT_NOTE
            else:
                par.add_run(SENT_NOTE)
            n += 1
        elif BARE in par.text and FULL not in par.text:
            n += replace_in_paragraph(par, BARE, FULL)
    d.save(str(FEMTECH))
    back = docx.Document(str(FEMTECH))
    text = "\n".join(p.text for p in back.paragraphs)
    bare_left = text.count(BARE) - text.count(FULL)
    print(f"  email document: {n} change(s); bare occurrences left: {bare_left}")
    return bare_left == 0


def fix_actionables() -> bool:
    shutil.copy2(ACTIONS, ACTIONS.with_name(ACTIONS.stem + " (backup 2026-09-29b).docx"))
    d = docx.Document(str(ACTIONS))
    hit = False
    for t in d.tables:
        for row in t.rows:
            if row.cells and "Accelerating FemTech: SENT" in row.cells[0].text:
                set_cell(row.cells[1], DONE_NOTE)
                hit = True
                break
        if hit:
            break
    if not hit:
        print("  Open Actionables: row not found")
        return False
    d.save(str(ACTIONS))
    back = docx.Document(str(ACTIONS))
    stale = any(
        "signature" in c.text and "Ruth Christianson\u201d" in c.text
        for t in back.tables for row in t.rows for c in row.cells
    )
    print(f"  Open Actionables: note rewritten; stale claim left: {stale}")
    return not stale


def fix_workbook() -> bool:
    wb = openpyxl.load_workbook(str(BOOK))
    ws = wb["Checklist"]
    changed = 0
    for row in ws.iter_rows(min_row=max(1, ws.max_row - 30), max_row=ws.max_row):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "was deliberately NOT corrected" in v:
                cell.value = "The sent email, corrected after the fact"
                changed += 1
            elif isinstance(v, str) and "because that is what the recipient received" in v:
                cell.value = (
                    "She edited the signature before sending, so the email went out as Ruth "
                    "Christianson-Monroy. I had left the draft reading 'Ruth Christianson' and written "
                    "that this was what the recipient received. The reasoning was right - a record of a "
                    "sent email must describe what actually went out - and the fact was wrong, so the "
                    "same reasoning now says to correct the file rather than preserve it. Corrected the "
                    "same day, in the document, in Open Actionables and here."
                )
                changed += 1
    wb.save(str(BOOK))
    print(f"  workbook: {changed} cell(s) corrected")
    return changed >= 1


if __name__ == "__main__":
    print("\n  CORRECTING THE RECORD\n")
    ok = all([fix_email(), fix_actionables(), fix_workbook()])
    sys.exit(0 if ok else 1)
