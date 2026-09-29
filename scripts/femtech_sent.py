"""Record that the Accelerating FemTech email went out, in both places.

Ruth, 29 September 2026: "The Accelerating FemTech email has been sent from
hello@selodia.app. Mark it done in Open Actionables and the funding log, dated
today."

Two documents, both edited in place with a dated backup first, because both are
hers and most of what is in them is history.
"""

import shutil
import sys
from copy import deepcopy
from pathlib import Path

import docx

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"H:\My Drive\Selodia App Project Master Folder")
LOG = ROOT / "Competitor analysis" / "Competitor analysis and funding watch.docx"
ACTIONS = ROOT / "Build Specs" / "Selodia - Open Actionables.docx"

LOG_ROW = (
    "29 Sep 2026",
    "Accelerating FemTech: EMAIL SENT to the Health Innovation Network South London team, asking when "
    "the early-stage programme reopens and whether there is a list to be on. Sent from hello@selodia.app, "
    "which is a Send As address on Gmail and both sends and receives. The programme itself is still closed "
    "with no published reopening date, so this is the only route open to us. Awaiting a reply.",
    "hin.southlondon@nhs.net",
)

DONE_WHAT = "Accelerating FemTech: SENT 29 September 2026"
DONE_NOTE = (
    "Done. Sent from hello@selodia.app to hin.southlondon@nhs.net, asking when the early-stage "
    "programme reopens and whether there is a list to be on. hello@selodia.app is set up as a Send As "
    "address in Gmail and both sends and receives, so the earlier caution about the from line is "
    "resolved and should not be raised again. The signature on the sent email reads \u201cRuth "
    "Christianson\u201d; it was left as sent because the document is now a record of what the recipient "
    "received. Her full name, for any reply or follow-up, is Ruth Christianson-Monroy. "
    "NEXT: a reply, or nothing. If nothing by the end of October, chase once."
)


def set_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for run in p.runs[1:]:
        run._element.getparent().remove(run._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def add_log_row() -> bool:
    if not LOG.exists():
        print(f"  MISSING  {LOG}")
        return False
    shutil.copy2(LOG, LOG.with_name(LOG.stem + " (backup 2026-09-29).docx"))
    d = docx.Document(str(LOG))
    table = None
    for t in d.tables:
        head = [c.text.strip().lower() for c in t.rows[0].cells]
        if head[:3] == ["date", "finding", "where"]:
            table = t
            break
    if table is None:
        print("  could not find the funding log table")
        return False
    if any("EMAIL SENT" in c.text for row in table.rows for c in row.cells):
        print("  funding log: already recorded")
        return True
    # NEWEST FIRST, which the document says above the table, so the new row goes
    # directly under the header rather than at the end.
    template = deepcopy(table.rows[1]._tr)
    table.rows[0]._tr.addnext(template)
    for cell, text in zip(table.rows[1].cells, LOG_ROW):
        set_cell(cell, text)
    d.save(str(LOG))
    back = docx.Document(str(LOG))
    ok = any("EMAIL SENT" in c.text for t in back.tables for row in t.rows for c in row.cells)
    print(f"  funding log: row added at the top, verified {ok}")
    return ok


def mark_actionable() -> bool:
    if not ACTIONS.exists():
        print(f"  MISSING  {ACTIONS}")
        return False
    shutil.copy2(ACTIONS, ACTIONS.with_name(ACTIONS.stem + " (backup 2026-09-29 before femtech sent).docx"))
    d = docx.Document(str(ACTIONS))
    hit = False
    for t in d.tables:
        for row in t.rows:
            if row.cells and "Accelerating FemTech" in row.cells[0].text and "NOT sent" in row.cells[0].text:
                set_cell(row.cells[0], DONE_WHAT)
                set_cell(row.cells[1], DONE_NOTE)
                hit = True
                break
        if hit:
            break
    if not hit:
        print("  Open Actionables: the FemTech row was not found")
        return False
    d.save(str(ACTIONS))
    back = docx.Document(str(ACTIONS))
    ok = any(DONE_WHAT in c.text for t in back.tables for row in t.rows for c in row.cells)
    still_pending = any(
        "NOT sent" in c.text for t in back.tables for row in t.rows for c in row.cells
    )
    print(f"  Open Actionables: marked sent, verified {ok}; stale 'NOT sent' left: {still_pending}")
    return ok and not still_pending


if __name__ == "__main__":
    print("\n  RECORDING THE SENT EMAIL\n")
    a = add_log_row()
    b = mark_actionable()
    sys.exit(0 if (a and b) else 1)
