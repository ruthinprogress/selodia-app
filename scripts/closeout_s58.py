"""Session 58 - the overnight run - into both sheets of the close-out workbook.

A SESSION OF ITS OWN, per WORKFLOW.md: "If the overnight work is itself
substantial, it gets its own close-out when she wakes, rather than being folded
silently into the previous day." Seven slices and eleven commits is substantial.
"""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

LABEL = "S58: 28Sep26  (overnight - Plans and onboarding spine)"

CHECKLIST = [
    ("Step 0: is the focus-state mechanism wired?", "Diagnosed", "2026-09-29",
     "Built since 9 September and wired to CHAT only. Onboarding never touched either column - the words fat_focus_state and muscle_focus_state appeared in the goals screen zero times. And the real default was deeper than the TypeScript: the columns were NOT NULL DEFAULT 'maintain', so the value was written at row creation before any code read it. All three accounts carried it, Ruth's included, set to maintenance while she trains for a muscle up."),

    ("Slice 1: goals reach the targets", "Done", "2026-09-29",
     "Onboarding is seven taps and writes the focus columns directly. The silent default is closed in all three places - the column default, asFocus on the server, asFocus in the app. Null now means not stated and not stated means NO target rather than a maintenance one. 'Keep things steady' is the only thing in the app that writes maintain, which is the whole point: a choice rather than a silence wearing a choice's clothes."),
    ("Today shows the target, six weeks late", "Fixed", "2026-09-29",
     "calorieTargetKcal and proteinTargetLabel have been computed in the Overview panel since 15 August and rendered NOWHERE. The app has known what a person's target was for six weeks and never told her. Stated, never counted down from. The first screenshot read 'Aiming for 1760 kcal - 83-99' because proteinTargetLabel returns a bare range, which nobody had noticed for the same reason."),
    ("Migrations, reversible", "Done", "2026-09-29",
     "Old values snapshotted into user_profile_focus_backup_20260928 with RLS on and no policy, which is the way back. Only rows carrying the default's exact signature were cleared - both at maintain, never updated since creation. The store account's deliberate reduce/increase was untouched. Five new tables, all additive, none of which anything read before tonight."),
    ("The checkbox announces its state", "Fixed", "2026-09-29",
     "One component, fourteen uses, including all three consent boxes and the eight that choose what goes in a GP report. A screen reader read the label perfectly and said nothing at all about whether the box was ticked, so a blind woman was agreeing to terms she could not confirm she had agreed to."),

    ("Slice 2: My Rules, enforced in code", "Done", "2026-09-29",
     "The spec has asked for this since the Movement brief, on the allergy gate's reasoning: a contraindicated movement is an injury risk and a prompt is a request. It runs on the way IN, so a plan that breaks a rule is never WRITTEN rather than merely never shown. An unconfirmed rule still excludes - the cautious side of an ambiguity, taken on purpose - and the screen says both that it is waiting and that it is being applied."),
    ("Rule capture confirms, never assumes", "Done", "2026-09-29",
     "A fifth type on the existing pending-save machinery: the offer lives in the database and the model only reports whether the answer was yes. A rule's offer says what agreeing will DO, because 'want me to keep that?' is not enough consent for something that changes what gets built for her from then on."),
    ("The letter upload is NOT built", "Blocked - needs Ruth", "2026-09-29",
     "Not for want of time. Accepting clinical letters needs the Play Data Safety form and the privacy policy changed first - both currently say files are not collected - and letters carry NHS numbers, addresses and clinicians' names. The spec already flagged it. Shipping an upload button ahead of those two documents would be a compliance problem, not a rough edge."),

    ("Slices 3 and 4: the Plans segments", "Done", "2026-09-29",
     "Week, Sessions, Skills, Rules, with goals above them on every segment. Week is not Sessions: a cadence and a thing to do, two objects, and a week row POINTS AT its session so there is never a second version of Tuesday. user_week has no completion column at all, so 'nothing is ever marked done or missed' is enforced by the schema rather than by everybody remembering."),
    ("THE CLIP LIBRARY CANNOT SHOW HER OWN LADDER", "Diagnosed", "2026-09-29",
     "Checked against all 923 clips: zero handstand, zero muscle up, zero scapular, zero dead hang, no bar dip - only rings. Every 'split' clip is a split SQUAT. So the muscle-up ladder is not seeded because it cannot be shown honestly, three that can be are, and the ten missing movements are a gap list in code where the next person writing a ladder will see it. Each line is a commission brief."),
    ("Skills have no timeframes", "Done", "2026-09-29",
     "A decision against her own prototype, on her own instruction. The prototype has 'MONTHS 3-6' and 'MONTHS 6-18'; the brief says Now/Next/Goal only. A month range on a skill is a deadline with better manners, and a woman still on the first rung in month seven has been handed a way to feel behind at something she took up for fun. There is no months column."),

    ("Slice 5: onboarding screens 2 to 7", "Done", "2026-09-29",
     "Skill and placement, life stage, what she already does, anything to steer around, how she would like to be led, and the first draft. Screen 5 saves NOTHING on purpose: a tap on 'an injury or condition' is not consent to anything, because the app has no idea yet what to exclude, and telling her 'noted' would leave her believing it knew about her shoulder."),
    ("The ninth life-stage branch", "Done", "2026-09-29",
     "Her addendum, and it is the case every app in this market gets wrong: a coil, or a hysterectomy with the ovaries kept, means no bleeding and a body that is still cycling. None of the other eight fit her and both answers she would pick are wrong. stageForReasoning returns null there, so anything wanting to say 'because you are post-menopausal' has to ask and be told no."),
    ("The cycle day answers to life stage", "Fixed", "2026-09-29",
     "The 45-day staleness ceiling protects a woman who stopped logging; it did nothing for one who had told us where she is, and a five-week-old period start on a post-menopausal profile still produced 'Cycle day 34'. Only a regular cycle gets one. Not sure and prefer-not-to-say get silence, which is the honest answer when the app does not know."),
    ("A bug caught by a screenshot", "Fixed", "2026-09-29",
     "The first-draft screen marked onboarding complete on mount, so the layout guard redirected her to Chat before the summary drew a single row. The screenshot showed a chat thread where her draft should have been - exactly what a real person finishing onboarding would have got. It advances when she taps Start, which is also when it is true."),
    ("Height and activity level nearly disappeared", "Fixed", "2026-09-29",
     "Neither was collected by any onboarding SCREEN before - the chat route extracted them from prose - so replacing those screens with taps would have removed the only path to both. Without height there is no BMR, no TDEE and no calorie target. Slice 1 would have fixed the target and Slice 5 would have removed its inputs."),
    ("Health context and nutrition targets no longer asked", "Outstanding", "2026-09-29",
     "The conversational screens that collected them are still in the app but no longer in the onboarding chain. A real gap, and hers to decide: fold them into the taps, or leave them to chat."),

    ("Slice 6: Guide me on Today", "Done", "2026-09-29",
     "'On the plan: ballet and gym', Guide me only, planned days only. A line, not a prompt, and never a tick: Today is not where anything is marked done. Somebody who does not do it finds the line gone tomorrow and nothing else changed - no carry-over, no count, no note that it did not happen."),
    ("Five notification strings written, none wired", "Proposed - needs her", "2026-09-29",
     "Each carries what it would say, when it would fire, and WHY it is worth interrupting somebody for - and if that last line is weak the notification should be cut. A sixth is written down as deliberately never existing: 'you have not logged in 3 days' is the app noticing she is absent and making it her problem, and no wording fixes it because the fact of the message is the judgement."),

    ("Slice 7: the red flags, built and switched off", "Done", "2026-09-29",
     "RED_FLAGS_LIVE = false, pending clinical advisor review. 18 flags: 6 for 999, 3 for 111, 9 for a GP. A switch rather than an unmerged branch so the code is on main, covered by its checks and cannot rot, and turning it on is one line and a clinical decision. Self-harm is deliberately NOT on it - the five-tier machine owns that, and a second dumber path would eventually disagree with the first."),
    ("The first matcher sent somebody to A&E for a question", "Fixed", "2026-09-29",
     "'Is chest pain always serious?' matched, because the abstract-question guard was a list of question openings and that sentence begins with none of them. A prefix list is the wrong shape: the ways to ask a question are unbounded and the ways to report something about yourself are not. Requiring a first-person reference is one rule instead of twenty."),
    ("The 12 health cases are listed, NOT run", "Queued", "2026-09-29",
     "Each needs a live turn to decide what the model says, and the output says so every time rather than letting a green tick imply otherwise. Ruth's sign-off was to BUILD, provisionally; running them against the live route is the next step and costs real money."),

    ("The flow map, retaken after the build", "Done", "2026-09-29",
     "Routes derived from the router tree rather than a hand-kept list, which is why it found six screens the old map had never heard of and stopped asking for three that no longer exist. Retaken again at the end so the six new onboarding screens are on it. Saved dated beside the 23 September one, never over it."),
    ("The clinical advisor", "Blocked - needs Ruth", "2026-09-29",
     "GBP 600-1,200 for a half day reading the boundaries, the red-flag list, the test set and the not-a-doctor wording. They would NOT review conversations or take clinical responsibility, and saying so belongs in any agreement because it is the first thing a clinician asks. This is the only thing blocking layer 2."),
]

DECISIONS = [
    ("One broken stand-in only catches one direction of wrongness", "2026-09-29",
     "Three check files ran every case twice, against the real implementation and a deliberately broken one, and all three came back with checks marked WEAK for the same reason - which I did not see until the third. The stand-in removed nothing, so every check about NOT over-removing passed against it trivially. A gate can be wrong in two directions, and in this codebase the eager half is the one that has actually bitten twice. Two stand-ins, and a check earns its place by failing against at least one; everything left over is a regression or sanity check and is counted separately rather than folded into the total."),

    ("When the exclusion list is unbounded and the inclusion test is small, invert it", "2026-09-29",
     "Telling 'I've got chest pain' from 'is chest pain always serious?' started as a list of question openings and failed on the first sentence tried. Every real report of a symptom contains a first-person reference and a question in the abstract contains none - one rule instead of twenty, and it never needs extending. The tell that a list is the wrong shape is that adding an item feels like it will be the last one and never is. Same shape as the allergy gate at one remove: a MATCH is not a REPORT, exactly as a MATCH was not a SUGGESTION."),

    ("A fabricated default lives in more places than the one you find it in", "2026-09-29",
     "The maintenance-target bug was reported as a line of TypeScript. Fixing that line would have changed nothing, because the column underneath was NOT NULL DEFAULT 'maintain' and the value was written at row creation. Three places - schema, server, app - and the schema was both deepest and least likely to be looked at, because a migration written months ago does not appear in a code review of today's bug. When a value appears where nobody chose it, check the schema before concluding you have found the cause."),

    ("An unconfirmed clinical rule still excludes", "2026-09-29",
     "The cautious side of an ambiguity, taken deliberately and recorded so it is not re-argued. If the app has heard 'my surgeon said no loaded squats' and has not yet had it confirmed, putting loaded squats in a session while waiting is the wrong way to be wrong. The screen shows it as awaiting confirmation AND says it is being applied meanwhile, because hiding the rule while keeping the effect is the worse half of both options."),

    ("A rule that removes movements trims the plan, it never refuses it", "2026-09-29",
     "Dropping a whole session because one movement broke a rule would punish her for the model's mistake and leave her with nothing. What was removed is named, with the rule that removed it - a session that quietly comes back two movements shorter teaches her the app is unreliable, and one that names the rule teaches her the rule is working, which is the only reason to have written it down."),

    ("A notification that cannot say why it is worth interrupting somebody should be cut", "2026-09-29",
     "Every reminder string carries a justification line, and the one that failed its own test is written down as deliberately never existing. 'You have not logged in 3 days' is every tracking app's standard message and the reason people delete them: a woman who stopped logging in a bad week gets told off by her phone for having had a bad week. There is no wording that fixes it, because the FACT of the message is the judgement."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    # Borrowed from an existing row rather than constructed, so the new rows
    # match whatever the sheet already looks like.
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 2  # one blank spacer row, per the ceremony
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value.startswith(("Done", "Fixed", "Diagnosed")):
                    cell.fill = copy(done_fill)
            r += 1
        return r - 1

    check_last = append(check, LABEL, CHECKLIST, 2)
    dec_last = append(dec, LABEL, DECISIONS, None)
    wb.save(BOOK)

    # READ IT BACK. "Do not report success from the fact that the write did not
    # raise" - the ceremony's own instruction, after the file being open in
    # Excel silently cost a close-out once.
    back = openpyxl.load_workbook(BOOK)
    print(f"  sheets: {back.sheetnames}")
    print(f"  Checklist: {back['Checklist'].max_row} rows, label at {check_last - len(CHECKLIST)}")
    print(f"  Decisions: {back['Decisions'].max_row} rows, label at {dec_last - len(DECISIONS)}")
    print(f"  label: {LABEL}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
