"""Session 59 - the overnight run of 29/30 September 2026 - into both sheets."""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

LABEL = "S59: 30Sep26  (overnight)"

CHECKLIST = [
    ("FUNDING STATUS", "Open", "2026-09-30",
     "Nothing moved overnight and nothing was due. FemTech email sent 29 Sept, awaiting a reply, chase once if nothing by end of October. Women in Innovation expected around 26 November. Both Monday tasks fire for the first time on 5 October - the search at 09:02 and the readiness check at 09:18."),

    ("HER CHAT NEVER LOST A MESSAGE", "Fixed", "2026-09-30",
     "loadThread read chat_messages with no limit, ordered ascending. PostgREST caps at 1000 rows and reports nothing, so the function's own error guard passed. At 1504 rows the app was handed the OLDEST thousand, ending 21 Sept 19:16, and the 504 since then were unreachable. Now newest-first with an explicit window of 300, reversed into reading order, plus one extra row to seed the discussion-card walk. Queries: count = 1504; the 1000th row by created_at = 2026-09-21 19:16:50; rows after it = 504."),

    ("check-row-ceiling.mjs", "Done", "2026-09-30",
     "Refuses the shape across the codebase. Its first run reported 42, which is a check nobody runs twice, and two of the three it called urgent were already safe - a drinks lookup bounded by an .in() of a few names, and a supersede window of a few seconds. It now understands that a filter bounds a result only when the thing it matches is small, and that a filter on `source` is not one of those, which is exactly what the broken read carried. 0 at the ceiling, 9 listed as worth a limit eventually."),

    ("I WAS WRONG ABOUT THE WAIST, AND IT IS CORRECTED", "Fixed", "2026-09-30",
     "It saved four seconds after she typed it, into personal_metrics: waist, 70, cm, raw_input 'Waist 70cm', created 20:24:39. I searched body_measurements, found no waist column, and reported 'nowhere' having proved only 'not here'. The peanut butter half was right and was re-checked with a query: no peanut butter in the dinner entry's text or any of its five items, and nothing at all written between 20:22 and 20:27."),

    ("Custom measurement tables", "Needs Ruth", "2026-09-30",
     "Created before checking what existed, then stopped at her instruction. They are EMPTY and NOT dropped. They duplicate personal_metrics, which has held her waist, thighs and calf since 27 August, and the Measurements screen derives its list of metrics from the readings themselves - which is why a calf appeared without anyone configuring one. Recommendation: keep personal_metrics, drop mine. Hers to decide."),

    ("A metric name is reconciled at the write", "Done", "2026-09-30",
     "The instruction to the model has asked it to reuse her existing metric names since the day it was written, and nothing enforced it. A model returning 'Waist' would have split her waist history in two, each half on its own line of a screen whose job is to show a measurement moving over time, with nothing looking wrong. Now reconciled where the row is built: her spelling wins, a genuinely new name stands as said. 10 checks, and the two about case fail against the old builder."),

    ("The end-to-end measurement test", "Queued", "2026-09-30",
     "'Waist 70cm', 'calf 36cm' and 'head 55cm' through the real parser need model calls, which cost money and were excluded tonight. The check prints the command to run rather than implying coverage it does not have."),

    ("An expanded Me entry says its body once", "Done", "2026-09-30",
     "The preview, the body and the grey line all showed the same sentence. The preview now stands down when the body appears, and the second line has to earn its place - measured against her three real cards: Nickel shares 0.63 of its vocabulary with its own why and is hidden, Seasonal Allergy shares 0.36 and is kept because it adds that the prescription is picked up and not yet used, Evening Skincare shares 0.00 and is kept. Labelled 'Latest' when shown. Edit and Talk this through untouched, as she instructed."),

    ("Two mistakes of mine inside that one", "Fixed", "2026-09-30",
     "I set the threshold at 0.65 from an estimate made by eye, which left Nickel's restatement on screen - the exact thing she reported. And my first fixtures used a `why` truncated to 160 characters by the query I happened to run, which made the Seasonal Allergy card look far more distinct than it is: its real text already names Grazax and the GP. A threshold tuned against truncated data is tuned against nothing."),

    ("Rename and voice", "Done", "2026-09-30",
     "'Skincare and allergies' became 'Allergies and intolerances'. The Nickel card is in second person - 'brings a rash up on your neck'. One row changed, snapshotted first, with the exact revert in scripts/snapshots/2026-09-29-nickel-before-rename.sql."),

    ("Multi-item entries render as a list", "Done", "2026-09-30",
     "A compact list with each item's own timing, rather than one blob. The reader is duplicated into mobile/src/lib because the mobile bundle has no path to app/lib; the duplication is narrow and guarded - check-me-items compares the two field lists and fails on drift, proved by removing a field from the copy."),

    ("Skincare card NOT converted to items", "Queued", "2026-09-30",
     "The script is written and the dry run is correct: three items, the contradictory 'retinol nightly' prose archived into history rather than deleted, same-day status flips collapsed. Not run, because tonight's rules forbid data patches on her account and this migration belongs to an earlier brief. It needs one word from her and one command."),

    ("NOTHING WAS SEEN WORKING", "Outstanding", "2026-09-30",
     "The web bundler died three times with 'Reached heap limit', never answering on its port. The machine has 15.9 GB with about 7.7 GB free while Chrome is open, and the bundle needs more old space than that. Lowering the ceiling to 6144 made it worse - it grew to 6.1 GB and died anyway - so 8192 is restored and the finding is written into start-web.mjs so nobody tries it again. Closing Chrome is usually enough. Everything tonight is covered by tests and unverified on screen."),

    ("Skills: the answer to her question", "Diagnosed", "2026-09-30",
     "No chat actions exist for skills - no add, no 'set where I am', no move a rung, no remove, no 'needs this first'. There is no server-side skills module at all; the only writer of user_skills and user_skill_rungs is the onboarding screen, and skills-view is read-only. Text-only rungs for clip-less moves WERE built, in skill-ladders.ts, and check-skill-clips refuses a text-only rung with no cue. 'No skills yet' is literal: both tables are empty, because the seeded ladders were removed as demo data at her instruction and nothing replaced them."),
]

DECISIONS = [
    ("A limit nobody set is a bug with a date on it", "2026-09-30",
     "PostgREST returns 1000 rows and says nothing about it. The read that broke her chat had no limit, an ascending order, and a guard that only checked for an error - so it failed silently, invisibly in review, and on a THRESHOLD rather than on a change. Nothing was edited on the day it broke; the table crossed a number, and a test written the week before would have passed. The rule that follows: a read of a table that grows with use must say how many rows it wants, and a filter counts as bounding only when the thing it matches is itself small."),

    ("Tune a threshold against the real text, not the text your query printed", "2026-09-30",
     "I built the 'does this line add anything' rule against a why truncated to 160 characters, because that is what my SELECT happened to show. The full text already named Grazax and the GP, so the card I was treating as the clear-cut keep-it case sat far closer to the line than I thought. The fixtures are now her actual wording, and the measured margins are written beside the threshold so the narrow one is visible to whoever changes it next."),

    ("An absence in one table is not an absence", "2026-09-30",
     "Recorded again from the day before, because it governed the whole night. I proved 'no waist column in body_measurements' and reported 'the reading was never saved'. Before reporting that something was not saved, find the path that WOULD have saved it and show it empty - or name which places were checked and say others may exist."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 2
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value.startswith(("Done", "Fixed", "Diagnosed")):
                    cell.fill = copy(done_fill)
            r += 1
        return r - 1

    append(check, LABEL, CHECKLIST, 2)
    append(dec, LABEL, DECISIONS, None)
    wb.save(BOOK)

    back = openpyxl.load_workbook(BOOK)
    print(f"  Checklist: {back['Checklist'].max_row} rows")
    print(f"  Decisions: {back['Decisions'].max_row} rows")
    print(f"  label: {LABEL}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
