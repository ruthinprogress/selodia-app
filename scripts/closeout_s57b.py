"""Append the second half of Session 57 to both sheets of the close-out workbook.

A second block under the same session rather than a new session: 57 is one day,
and a workbook that invents S57a and S57b would make the run of session numbers
stop meaning anything.
"""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

CHECKLIST = [
    ("The standing rule did not hold, and why", "Done", "2026-09-28",
     "Checked rather than assumed: items 2-7 arrived in the same message as the correction and were not skipped. The real gap is the rule's last line - it says WHEN to close out and nothing about what happens next, so a close-out reads as a full stop. Now: a close-out is a checkpoint, the session ends when she says so, and work arriving mid-session joins the queue."),
    ("Voice on the old path, text on the new", "Done", "2026-09-28",
     "The switch is a decision per surface now. Seven seconds of typed reply is a pause with a typing indicator; seven seconds of spoken reply is silence on a phone call."),
    ("24 September: five duplicate weight rows removed", "Done", "2026-09-28",
     "All six were written on the 26th, back-dated, over 28 seconds - 58.0, then 58.0 with body fat, then 55.5 complete, then three more copies. A correction typed in stages that wrote a row each time. One 55.5 kept, five backed up to a new body_measurements_removed table."),

    ("Latency: 3.3s is not reachable with two sequential calls", "Diagnosed", "2026-09-28",
     "Measured floors: classify ~1.9s with the reply asked as a fallback, writer ~2s at best, and the old path WAS the classify call. Running the two at the same time is the only route there, and on a turn with nothing to report the writer needs nothing the classifier produces. Not built - a real architectural change, and it spends a wasted call on the minority of turns that do have something to report. Her question, in progress.md."),
    ("Waiting list: where, how many, is it live", "Done", "2026-09-28",
     "The `waitlist` table in Supabase, written by a Server Action on the landing page. NOT Jotform. One person: Fiona, 2 September. Proved live in a browser - a curl POST returns 200 and writes nothing, because a Server Action needs the browser's own machinery, and that was one step from being reported as a bug."),
    ("Why she could not see the waiting list", "Done", "2026-09-28",
     "RLS grants INSERT to anyone and NO SELECT at all - right for a list strangers add themselves to, and it locks her out too. Twenty-six days of collecting with no route to the contents. scripts/waitlist.mjs, --csv for a sheet."),

    ("Running costs and pricing", "Done", "2026-09-28",
     "Computed from measured token counts. A chat turn 1.77c, a logging turn 2.01c, and caching is doing half that work. About GBP 2.09 a user a month blended, and FLAT from 100 to 10,000 - no economy of scale to wait for. Recommendation GBP 6.99/month, GBP 59.99/year. The biggest item on the page is not a cost: the ElevenLabs grant ends September 2027."),
    ("A heavy user costs more than they would pay", "Unresolved - watch it", "2026-09-28",
     "GBP 11.88 in models against GBP 5.94 of revenue at GBP 6.99. Today that is one person and she built it. Instrument cost per user before designing against it - every Anthropic response already carries its usage."),

    ("DPIA drafted", "Done", "2026-09-28",
     "docs/dpia.md and the new Legal folder on Drive, with the live privacy policy and terms filed beside it. 39 tables inventoried, six risks with residuals. Four things named as wave-one blockers, one of them new: the international transfer mechanism for the two US processors. Whether a DPO is required I could not establish and did not guess."),
    ("Beta agreement v1.1, and acceptance in the app", "Done", "2026-09-28",
     "Her final text. beta_members and beta_agreement_acceptances - two facts, not one boolean, because 'granted but not yet accepted' is a real state clause 13 creates. Versioned and append-only like consent_records."),
    ("Three placeholders block the agreement being shown", "Outstanding - needs her", "2026-09-28",
     "The version date, the registered address and the company number. All follow incorporation."),

    ("Beta feedback built and shipped", "Done", "2026-09-28",
     "A wave-one blocker. Top of More for beta accounts, one tap from anywhere via the seed mark, which now records the route it was pressed FROM. NOTHING IS REQUIRED to press Send - a single tapped feeling is a complete submission, and the Send button is never disabled. History with sent/read/fixed underneath, because somebody who reports into silence stops reporting."),
    ("Screenshots: gallery only, private bucket", "Done", "2026-09-28",
     "Under the sender's own user id, enforced by the storage policy rather than the client choosing a path. Capturing the previous screen needs a native view-shot module that is not in this build - a whole new binary to ship a form."),
    ("scripts/beta-feedback.mjs", "Done", "2026-09-28",
     "Lists every submission and sets its status. The status is a promise to the tester: marking something fixed when nothing was fixed is lying to the person who took the trouble."),
    ("The three old Jotform forms", "Outstanding - hers", "2026-09-28",
     "Two Unflump beta feedback forms and Beta Tester Coffee Chat Notes. Replaced by the in-app screen; archive them when she is ready."),

    ("Drinks are measured, not estimated", "Done", "2026-09-28",
     "Half a lager at 180 kcal and 6 g protein where CoFID makes a 284 ml half 68 and 0.9 - protein seven times out. At the write, where the text logger and image parser both arrive, with the parent totals recomputed. 14 checks, including everything it must NOT touch."),
    ("hello@selodia.app does not receive mail", "Outstanding - needs her", "2026-09-28",
     "Already printed on the landing page and the in-app support screen, and a mailbox is separate from the domain being on Vercel. Google Workspace ~GBP 6/user/month or Fastmail ~GBP 4; either needs two DNS records. DNS is a domain change and hers - the exact records are waiting on which provider."),
    ("Apple Small Business Programme", "Outstanding - reminder owed", "2026-09-28",
     "The day enrolment 3N9H5LB49A is approved. 30% to 15% from year one; a form, not a negotiation."),
    ("Session close-out ceremony", "Done", "2026-09-28",
     "Second close-out of the same session. Build log appended to the existing S57 entry rather than opening a second one; article log, decision patterns, spec, beta checklist and Open Actionables all updated; five documents synced and verified."),
]

DECISIONS = [
    ("A permission set correctly in one direction can be a gap in the other", "2026-09-28",
     "The waitlist table allows INSERT to anyone and grants no SELECT, which is right for a list strangers add their email to - a stranger who can read it has everybody else's address. What nobody followed through is that the same rule locks out the person the list is FOR. Twenty-six days of sign-ups with no route to them. A security decision is usually made while thinking about the attacker, and the legitimate reader is a second thought that never arrives."),
    ("Find out how the thing is actually invoked before concluding it is broken", "2026-09-28",
     "The waiting-list form was tested with a command-line POST. It returned 200 and wrote nothing, which is exactly what a broken form looks like, and it was one step from being written up as a bug - a Server Action needs a header and an encoding the browser supplies. Same mistake twice in two days: a latency probe whose fixture described a turn that could not happen, and this. When a test of a live system fails, ask first whether the test invoked it the way anything real does."),
    ("A rule that says when to do something and not what happens afterwards fails at that seam", "2026-09-28",
     "The standing rule held for a full day and then did not. The sentence at fault said to close out when everything was done and stopped there, so closing out read as ending. Nothing was disobeyed: the rule was complete about the moment it described and silent about the moment after, and silence in an instruction is filled by whatever the reader already assumes. Worth applying to any process document with a final step - the question is not 'is this step clear' but 'what does somebody do the second after finishing it'."),
    ("Design for the person who will not fill anything in", "2026-09-28",
     "The beta feedback screen replaced three external forms that had collected nothing. One instruction shaped it: nothing is required to press Send. A single tapped feeling with no words is a complete report. That inverts the usual shape and is right for the reason the old forms failed - the moment worth capturing is when somebody is irritated and holding a toddler, and any friction loses it. What makes it work rather than merely permissive is the other half: every submission shows sent, read or fixed, because somebody who reports into silence stops reporting."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 1
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value in ("Done", "Diagnosed", "Fixed"):
                    cell.fill = copy(done_fill)
            r += 1

    append(check, "S57 continued: 28Sep26  (afternoon queue)", CHECKLIST, 2)
    append(dec, "S57 continued: 28Sep26", DECISIONS, None)
    wb.save(BOOK)
    print(f"Checklist {check.max_row} rows, Decisions {dec.max_row} rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
