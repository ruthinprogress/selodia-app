"""Append Session 56 to both sheets of the close-out workbook.

Two continuous sheets, appended to. Never a tab per session - that decision is in
WORKFLOW and the whole value of the workbook is that one scroll covers the project.
"""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

CHECKLIST = [
    # ── The chat audit and rebuild (her item 10) ───────────────────────────────
    ("Audit: why the in-app chat is worse than the model", "Done", "2026-09-27",
     "Not the model - chat, voice and roundup all run Sonnet 5. Four things around it: the reply was a field in a 49-field tool, ~16,000 tokens of instruction in front of it, six of eight data blocks passed as unrounded prose, and up to eight notes appended after the model had finished."),
    ("Offline baseline comparison, no app changes", "Done", "2026-09-27",
     "Replaced her item 11 (founder mode) at her instruction. A 277-token baseline beat the live app on all four of the week's failures. That is what the rebuild started from rather than trimming GENERAL_CONDUCT."),
    ("Prompt rebuilt from the baseline upwards", "Done", "2026-09-27",
     "770 tokens typed, 884 spoken, 838 roundup, against 12,593 for GENERAL_CONDUCT + CAPABILITIES. Six rules earned their way back, each naming the test that fails without it. Incident notes and statistics moved to docs/chat-prompt-history.md."),
    ("All numbers computed in code, rounded to the screen", "Done", "2026-09-27",
     "app/lib/turn-facts.ts, all eight blocks. The 1.4 had a mechanism: the measurement block interpolated an unrounded 55.58 where her screen showed 55.6, and the model did the subtraction itself."),
    ("An empty log says so, explicitly", "Done", "2026-09-27",
     "This is what stopped the invented hard session. A list somebody has to NOTICE is empty is what produced it. Same rule now in the roundup card."),
    ("Reply out of the tool, written after the saves", "Built - SWITCH OFF", "2026-09-27",
     "app/lib/chat-path.ts. REPLY_WRITTEN_AFTER_THE_SAVES is false: every turn still runs the old path. A failure in the new one falls back to the old reply rather than costing a turn. She reads the before-and-after first, then uses it for a day."),
    ("Nothing appended to a finished reply", "Built - SWITCH OFF", "2026-09-27",
     "The eight notes go IN as labelled facts before anything is written. Tested on the coffee turn, which under the old path said 'logged fine' and then 'it looks like that entry didn't save': one coherent message now. The Almanac offer stays outside, recomputed, because it must be asked exactly once."),
    ("save-honesty: item 6's fix had only reached one branch", "Done", "2026-09-27",
     "Found while testing the new path. The partial-miss branch - a weight saves while a waist does not, which is the COMMON case - still said 'would you mind re-entering it so we can make sure it's properly logged for you'. Her exact complaint."),

    # ── The tests, which were the day's real finding ───────────────────────────
    ("Two eval checks passed while testing nothing", "Done", "2026-09-27",
     "Written through a shell heredoc, which turned the word-boundary escape into a literal backspace. Neither regex could match. Both PASSED on a reply that said 'Got it, 56.9 kg logged'."),
    ("Both eval columns were flying blind", "Done", "2026-09-27",
     "The split between what goes to both columns and what goes only to the new one was made on the cases and not in the two functions reading them, which still looked for a field no case defines. Every reply in that run was written without being told what happened to her data."),
    ("A case whose five checks could not fail", "Done", "2026-09-27",
     "All five were negatives, so all five passed on the reply 'What's on your mind?'. Nothing was there to be wrong."),
    ("Guard: the suite refuses to run until the checks can fail", "Done", "2026-09-27",
     "Every check runs against an empty reply before a token is spent; a case scoring full marks on nothing stops the run, as does any check containing a control character. Found two more unfailable cases on its first run and a mangled regex on its second."),
    ("Five checks were testing a wording, not a behaviour", "Done", "2026-09-27",
     "One demanded the literal words '3 readings' and failed 'Weigh-ins: 3'. One demanded 'saved' and failed \"it's not in your log\". Each fixed and commented with the reply that exposed it."),
    ("Full test set: before 31/35, after 34/35", "Done", "2026-09-27",
     "And the honest reading is not that the new prompt wins the score. Once BOTH prompts get correct rounded figures most of the gap closes. It was never mainly a prompt problem."),

    # ── Item 9, layout half ───────────────────────────────────────────────────
    ("Weekly roundup: its own card, rows not prose", "Done", "2026-09-27",
     "Rows computed in app/lib/roundup-figures.ts from her own stored rows, never parsed out of the reply - parsing would have produced tidy rows of the same wrong figures. Verified against the exact week she was shown: 17 checks."),
    ("The roundup had never read her steps", "Done", "2026-09-27",
     "Not in any version of that route. The week it called 'almost no movement to speak of' ran 2,262 to 9,820 steps a day. The sessions were in the prompt; the steps were not read at all, so no instruction about fairness could have helped."),
    ("Where 'the thread running through this week is permission' came from", "Diagnosed", "2026-09-27",
     "Step 4 of the roundup prompt's numbered ORDER list: 'one thematic observation drawn across the week'. Run three times it produced unevenness, patchiness and incompleteness. It was asked for a theme and gave one."),
    ("chat_messages.kind and .meta", "Done", "2026-09-27",
     "Both nullable, migration applied and verified. Not `source`: Chat reads .eq('source','chat'), so 'roundup' there would have removed the roundup from the thread it belongs in."),
    ("Weekly roundup: content half", "Queued", "2026-09-27",
     "Her item 9's other half. The prose still comes from a prompt that asks for a theme. Not started, and it is the next thing on the roundup."),

    # ── Item 12 ───────────────────────────────────────────────────────────────
    ("Build log missing sessions 45 to 53", "Done", "2026-09-27",
     "Nine sessions, 16 to 24 September. Every one has a full block in the workbook, so the sessions WERE closed out - only this artefact was skipped. Backfilled from the workbook and the commit history, each entry saying so."),
    ("Why the build log stopped", "Done", "2026-09-27",
     "WORKFLOW called it 'appended automatically at step 5 of the automated close-out'. It never was; no script for it existed in the repository. And that sentence was written while fixing the OPPOSITE fault in the same step. Corrected, with the story kept visible."),
    ("Article log missing 25, 26 and 27 September", "Done", "2026-09-27",
     "Nine entries written, including her line about five of six incidents. Word copy regenerated."),
    ("The build log's Contents list had stopped at session 44", "Done", "2026-09-27",
     "Eleven sessions in the document and not in its own index. It is now rebuilt from the document's own headers on every append, so there is no second list to keep in step."),
    ("Sessions 54 and 55 used an em dash where every other entry uses a middle dot", "Done", "2026-09-27",
     "Against the document's own convention, WORKFLOW's template, and her writing rule. Normalised. It also broke the first version of the contents rebuild, which silently omitted the two newest sessions."),
    ("scripts/build_log_append.py", "Done", "2026-09-27",
     "Backs the file up first, refuses a session already present, rebuilds the contents, --before N for a backfill. It does not write the prose, deliberately."),
    ("closeout_check.py now checks the written logs and BLOCKS", "Done", "2026-09-27",
     "This session missing from the build log, any gap in the run of session numbers, no article entry for today, a stale article .docx, no workbook rows for today. A warning at the end of a long report is a thing to scroll past."),
    ("Location and start time for sessions 45 to 53", "Unresolved", "2026-09-27",
     "Never recorded anywhere and not guessed at. The nine backfilled headers say 'location and start time not recorded', the same convention sessions 38 and 39 already use."),

    # ── Carried, and hers ─────────────────────────────────────────────────────
    ("Supabase auth config: redirect URL and email template", "Outstanding - needs her", "2026-09-27",
     "Needs a personal access token (sbp_), not the service role key: the service key has authority over DATA, not CONFIGURATION. SUPABASE_ACCESS_TOKEN=sbp_xxx node scripts/supabase-auth-config.mjs. She is creating the token at her next laptop session."),
    ("Password reset tested on a real phone", "Outstanding - hers", "2026-09-27",
     "The app half is built. Until the Supabase side is set and this is tested, a beta tester who forgets their password is still stuck."),
    ("Turn the new chat path on", "Outstanding - hers", "2026-09-27",
     "She reads the before-and-after replies first, then says. One line, one commit, one web redeploy. The old path comes out after a day of use."),
    ("Plan-as-session routing", "Queued", "2026-09-27",
     "Carried from her priority order. Not started."),
    ("Three scale marks (weight, body fat, muscle)", "Queued", "2026-09-27",
     "Carried from S55. She approved drawing them in the same family as the tape marks. Not started."),
    ("Botanical drawings: screenshot where each appears", "Queued", "2026-09-27",
     "Carried from S55. Needed the bundler, which now works, so this is unblocked for the first time."),
    ("Repeated-phrase probe over a 10-minute conversation", "Queued", "2026-09-27",
     "Carried from S55. probe-reply-variety.mjs measures openers across stored replies; her ask is for repeats WITHIN one conversation."),
    ("Almanac > Me restructure to a record list", "Queued", "2026-09-27",
     "Carried. The largest remaining piece of her item 6 and not started."),
    ("ElevenLabs key rotation, and the access token in their records", "Outstanding", "2026-09-27",
     "Carried, untouched. The only carried item with real exposure."),
    ("Expo web bundler degradation", "Fixed", "2026-09-27",
     "Carried from S54 and S55 as unresolved both times. Two dev servers shared one Metro cache: 54,817 failed writes, nothing ever cached, every bundle cold. Warm bundle 27s. docs/bundler.md."),
    ("Session close-out ceremony", "Done", "2026-09-27",
     "Run overnight while she slept, at her instruction. Workbook, build log, article log, decision patterns, spec, beta checklist and open actionables all written; documents synced and verified; tree clean and pushed."),
]

DECISIONS = [
    ("When every rule you add makes the thing worse, the rule was never the fix", "2026-09-27",
     "Five of six chat incidents had been answered with a prompt rule when the cause was a missing field, a missing fact, an unrounded number or a vendor setting. The rule never fixed the incident; it made the prompt longer and gave every genuine instruction one more thing to compete with. So the rebuild inverted the burden of proof: start from a minimal baseline, add a rule back only when a test fails without it. Six earned a place out of 11,666 tokens of predecessors."),
    ("A test that cannot fail is worse than no test, because it reports success", "2026-09-27",
     "Twice in one evening the suite printed PASS on a reply that plainly failed - once from a backslash a shell turned into a control character, once from a case whose checks were all negatives and so were satisfied by an empty answer. Care was being taken both times, so resolving to be more careful was not available as a fix. The guard went at the write: the suite now refuses to run until every case fails something on an empty reply."),
    ("One message, one author", "2026-09-27",
     "Up to eight notes could be appended to a reply the model had already finished. That is the whole mechanism behind the reply that said the coffees were 'logged fine' and then that the entry 'didn't save'. No prompt rule could have prevented it, because the model had stopped writing before the contradiction was added. If the app knows something the reader needs, the model needs it before it starts, not the reader afterwards."),
    ("A layout cannot rescue a figure", "2026-09-27",
     "The obvious build for the roundup's UI pass was to split its prose into rows. It would have produced three well-spaced rows saying 1,400 kcal over a week averaging 1,222, and ten readings where there were three. The rows are computed from her own data instead. And building it found why one figure was wrong: the roundup had never read her steps, so 'almost no movement' was the honest conclusion from what it had been handed."),
    ("A step described in the documentation as automatic is a step nobody checks", "2026-09-27",
     "Nine sessions missing from the build log, with the workbook complete for every one of those dates. The cause is one sentence calling the step automatic when no script for it had ever existed - written, as it happens, while fixing the opposite fault in the same step. A chore that reads as already done is more dangerous than one that nags, because the nagging one gets checked. The artefacts written by scripts had no gaps; both written by hand had drifted."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 1
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value in ("Done", "Diagnosed", "Fixed"):
                    cell.fill = copy(done_fill)
            r += 1

    append(check, "S56: 27Sep26  (08:30 start, mobile out with Felix then overnight on the laptop)", CHECKLIST, 2)
    append(dec, "S56: 27Sep26", DECISIONS, None)
    wb.save(BOOK)
    print(f"Checklist {check.max_row} rows, Decisions {dec.max_row} rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
