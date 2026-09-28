"""Append Session 57's fourth queue to both sheets of the close-out workbook."""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

CHECKLIST = [
    ("The spoken reply streams, behind a switch", "Done", "2026-09-28",
     "REPLY_STREAMS_TO_VOICE, one line to turn off. The route's shape does not change: the adapter calls the pipeline in process, so the words go into a sink it already holds rather than the reply becoming a streaming response - which would have meant restructuring the offer line, the allergy gate and the one-reply-per-turn insert, all load-bearing and none of it about streaming."),
    ("Nothing is spoken until three things are true", "Done", "2026-09-28",
     "The turn is ordinary, the reply written ahead had nothing to report, and the allergy gate CANNOT ARM on that account - decided before a word exists, because the gate short-circuits for anybody with no food exclusions. So words already spoken can never need taking back. Everybody else waits, which is slower and is the only honest answer."),
    ("The 1.9s estimate was wrong, and why", "Diagnosed", "2026-09-28",
     "Waiting for a full stop measured 124ms on the median turn. The estimate came from a real gap - saves confirm 1.9s before the writer's last token - and assumed the writing was SPREAD across it. It is not; almost all of it is the model producing nothing yet. A duration between two events says how long something took, not how the work inside it is distributed."),
    ("Cutting at a clause instead of a sentence", "Done", "2026-09-28",
     "A comma, semicolon, colon or dash with at least 35 characters in front of it. 680ms median, 2,535ms best, against 124ms and 806ms. Sixty characters was the first floor and was too high - a real reply's opening clause is 43 - and the check caught that rather than a measurement afterwards."),

    ("THE REBUILT REPLY WAS NEVER GATED", "Fixed", "2026-09-28",
     "runAllergyGate runs on the OLD path's reply, and turnIsOrdinary only asks whether THAT one passed. The rebuilt path then stores a completely different sentence. For anybody with a declared food allergy the gate was examining a draft that was thrown away - worse than no gate, because the turn is recorded as gated. Survived a typecheck, a lint and a day of use. Nobody was affected: the only account with allergies has two that are not food, which is luck rather than design."),
    ("scripts/check-gated-paths.mjs", "Done", "2026-09-28",
     "Reads the route and fails if model-written words reach the reply without a gate in front of them. Proved against the shipped code, where it reports this exact bug, and against a copy with the new call deleted."),

    ("A backup of her chat messages was public for a day", "Fixed", "2026-09-28",
     "chat_messages_duplicates_removed, created 27 September before deleting the duplicates. Backing up first was right. What was missed is that a table in `public` is served by PostgREST, and without RLS the ANON KEY reads it - the key inside the mobile app and the landing page's own JavaScript. Found by the security advisor, opened to answer a different question."),
    ("scripts/check-rls.mjs", "Done", "2026-09-28",
     "Tries every table PostgREST serves with the anon key, because the question is not 'is RLS enabled' but 'can that key read it'. 43 tables, none returns a row. Proved able to fail: a table with one harmless row was created open, the check named it and exited 1, and it was dropped."),

    ("Founding rate: kept for good, or capped", "Answered - her decision", "2026-09-28",
     "Section 8 of the pricing document, with the arithmetic for all three shapes. 50 founders at GBP 6 against GBP 9 forgoes GBP 1,530 a year - and that is the wrong number, because those 50 would not all have subscribed at GBP 9. Recommendation: kept for good, capped at 100 places. NOT the 12-month cap, which is the cheapest and the one that spends trust, and needs more machinery rather than less."),
    ("Leaked-password protection: where, and what it costs", "Answered - needs her", "2026-09-28",
     "Authentication > Sign In / Providers > Email > Prevent use of leaked passwords. Supabase's advisor confirms it is disabled. IT NEEDS THE PRO PLAN at $25/month. So it is a GBP 20-a-month decision rather than a toggle. I called it a wave-one blocker in the DPIA and that was too strong - it is a real improvement and not the line between lawful and unlawful. The better argument for Pro is 7-day log retention, which would have answered two questions today that the free tier lost."),
    ("ICO moved to a wave-zero blocker", "Done", "2026-09-28",
     "Her decision: registers just before inviting Nikki and Carol rather than at launch. Recorded in the spec, the DPIA and the Open Actionables. It is the ONE item the wave-zero reasoning does not reach - everything else can be carried by hand for two people she can telephone, and this is owed to a regulator."),
    ("BotanicalMark", "Done", "2026-09-28",
     "Already gone, deleted earlier in Session 57. Nothing named botanical remains anywhere in the repository. Checked rather than assumed."),
]

DECISIONS = [
    ("A safety check that runs on a value something else then replaces is worse than no check", "2026-09-28",
     "The allergy gate examined the old path's reply; the rebuilt path then wrote a different sentence and stored it. The turn is recorded as gated, which is the part that makes it worse than nothing. Nothing about it is visible in the shape of the code - both branches assign to the same variable and the check sits above both. The general form: when a value is validated and then REASSIGNED, the validation belongs to the reassignment rather than to the variable. The question that finds it: of everything that can end up in this variable, which has been through the check? If the answer is 'the first one', the check is in the wrong place."),
    ("A table created in a hurry, for a good reason, skips every habit a real table gets", "2026-09-28",
     "The backup table was made carefully, with a comment explaining why, in the same minute as the deletion it protected - by somebody thinking about the deletion. A feature gets a migration review, an RLS policy and a line in the spec. A backup table gets none of them, and that is precisely what made it dangerous. The rule: a table is a table, and if it is created in a schema the API serves it gets row-level security in the same statement that creates it."),
    ("A gap in a timeline is not automatically an opportunity", "2026-09-28",
     "Streaming was estimated at 1.9 seconds from a real measurement - the saves confirm 1.9s before the writer's last token - and measured 124ms. The gap is real; what the estimate assumed is that the work was spread across it, when almost all of it is the model producing nothing yet. A duration between two events tells you how long something took and not how the work inside it is distributed, and a plan that depends on the distribution needs it measured. The tell was available and unasked: nobody had measured when the FIRST sentence appeared, only when the last one did."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 1
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value.startswith(
                    ("Done", "Fixed", "Answered", "Decided", "Diagnosed")
                ):
                    cell.fill = copy(done_fill)
            r += 1

    append(check, "S57 continued: 28Sep26  (fourth queue)", CHECKLIST, 2)
    append(dec, "S57 continued: 28Sep26  (fourth queue)", DECISIONS, None)
    wb.save(BOOK)
    print(f"Checklist {check.max_row} rows, Decisions {dec.max_row} rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
