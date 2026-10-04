"""Append one session to both sheets of the close-out workbook.

    python scripts/closeout_append.py session.json

WHY THIS REPLACES closeout_s55.py ... closeout_s58.py AND THE REST. There are
eight of those, one per session, each a near-copy of the last with a different
list of strings in it. Ruth, 4 October 2026, about the documents the same habit
produced in her Drive folder: "there's tons of backups across multiple files and
it's very cluttered." The scripts had it too - a new file every time instead of
one that takes the content as data.

A new file per session also means the APPEND LOGIC was copied eight times. Any
fix to it - the spacer row, the fill borrowed from an existing cell, the read-back
- had to be made in the copy somebody happened to open.

THE JSON IS THE SESSION, and nothing else is:

    {
      "label": "S61: 1Oct26",
      "checklist": [["Task", "Status", "2026-10-01", "Notes"], ...],
      "decisions": [["Decision", "2026-10-01", "Reasoning"], ...]
    }

A FUNDING LINE IS REQUIRED, because it is the one that goes missing. The rule is
"one line on funding in every close-out", and a rule that depends on remembering
is a rule that lapses - it already had. This refuses to write a checklist with no
funding row rather than leaving it to the next person's attention.
"""

import json
import sys
from copy import copy
from pathlib import Path

import openpyxl

BOOK = Path(
    r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"
)


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2

    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    label = data["label"]
    checklist = [tuple(r) for r in data["checklist"]]
    decisions = [tuple(r) for r in data["decisions"]]

    if not any("funding" in str(r[0]).lower() for r in checklist):
        raise SystemExit(
            "  No funding line in the checklist, so nothing was written.\n"
            "  Every close-out carries one, even when the answer is 'no movement'."
        )

    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]

    # Borrowed from an existing row rather than constructed, so new rows match
    # whatever the sheet already looks like.
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, rows, status_col):
        r = ws.max_row + 2  # one blank spacer row, per the ceremony
        ws.cell(r, 1, label).font = copy(header_font)
        start = r
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and str(value).startswith(
                    ("Done", "Fixed", "Diagnosed", "Built")
                ):
                    cell.fill = copy(done_fill)
            r += 1
        return start

    check_at = append(check, checklist, 2)
    dec_at = append(dec, decisions, None)
    wb.save(BOOK)

    # READ IT BACK. "Do not report success from the fact that the write did not
    # raise" - the ceremony's own instruction, after the file being open in Excel
    # silently cost a close-out once.
    back = openpyxl.load_workbook(BOOK, read_only=True)
    ck, dk = back["Checklist"], back["Decisions"]
    written = ck.cell(check_at, 1).value if ck.max_row >= check_at else None
    if written != label:
        raise SystemExit(f"  Read-back failed: expected {label!r} at row {check_at}, found {written!r}")
    print(f"  {label}")
    print(f"  Checklist: {len(checklist)} row(s) from {check_at + 1}, sheet now {ck.max_row}")
    print(f"  Decisions: {len(decisions)} row(s) from {dec_at + 1}, sheet now {dk.max_row}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
