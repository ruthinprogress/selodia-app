"""Session 58's afternoon - the drag - into both sheets of the close-out workbook."""

import sys
from copy import copy

import openpyxl

BOOK = r"H:\My Drive\Selodia App Project Master Folder\Build Specs\Selodia-Session-Closeouts.xlsx"

LABEL = "S58 continued: 29Sep26  (the drag)"

CHECKLIST = [
    ("Plans polished against her design", "Done", "2026-09-29",
     "Goals as cards with a Goal screen behind them showing what the goal drives, computed from her profile with the same functions Today uses rather than stored. All four segment labels in full - measured, not guessed: tabs are 80px and the longest label is 50px, so the padding was eating room that was always there. The week as seven real days with pills, a '+' per day, Anytime in two columns and Walking reading the phone's steps. Logging from a card and for the whole week, everything up to today arriving ticked."),

    ("The seed mark was a disc parked over content", "Fixed", "2026-09-29",
     "Her report: it sat over '1.5 hrs' on the Ballet card. The 25 September fix gave the mark something to sit ON, which stopped content running THROUGH the seeds and still hid whatever was underneath - from the reading side, the same complaint. A full-width band in the page colour now, so content passes behind a clean edge instead of colliding with a shape."),

    ("'Let me lead' could not do what the brief said", "Fixed", "2026-09-29",
     "'Everything in Anytime unless dragged' was impossible as built: the day rows were hidden entirely in that mode, so there was nowhere to drag to and moving a card changed nothing on screen. The deeper fault is that the days on a plan are a SUGGESTION written for everybody, and nothing distinguished that from a choice she made. user_week.days_chosen_at now records it, stamped at the write, and placedOnDays() is the single answer both the week and the log sheet ask - so the sheet can no longer offer Thursday for something the week is showing in Anytime."),

    ("DRAG DID NOT EXIST, AND I HAD SAID IT DID", "Fixed", "2026-09-29",
     "Her report: 'long-pressing and dragging does nothing at all', and she was exactly right - there was no drag gesture in the code. I had built the hold as a shortcut to a chooser sheet, documented the trade-off at the top of the file, and then described it to her as 'hold a card and move it', which reads as drag. She spent a message diagnosing a feature that had never been built."),

    ("Her first question answered before anything was built", "Done", "2026-09-29",
     "Whether it needed a native module missing from her 21 September build, because that decides everything else. It does not: gesture-handler and reanimated have been dependencies since the first mobile scaffold on 9 August, so they are compiled into her binary. GestureHandlerRootView - absent until 21 September, which is why no gesture in this app had ever fired - is JavaScript and already reached her over the air. No reinstall, and nothing changes for her data or settings."),

    ("The hold was also fighting the ScrollView", "Fixed", "2026-09-29",
     "Her other suspicion, and correct. Pressable's onLongPress lives in React Native's JS responder system, where a parent ScrollView can claim the responder as soon as a finger moves - and claiming it CANCELS the pending long press. 'Press and drag' cancelled itself before the timer fired, which is why tapping worked and holding did not. Gesture.Pan().activateAfterLongPress() is built for exactly this."),

    ("Drag built, and confirmed by her on the phone", "Done", "2026-09-29",
     "Drop zones measured with measureInWindow at the moment she lifts a card, so nothing needs the scroll offset. The card dims where it was, a preview follows her finger, the destination is outlined, the drop saves at once, and a drop that lands nowhere puts it back rather than guessing. Her words: 'Drag works on my phone.'"),

    ("'Move to…' is permanent and is the primary route", "Done", "2026-09-29",
     "Her instruction, and her reasoning is the right one: a drag cannot be performed with a screen reader, one-handed, or by anyone whose hands are having a bad day - which for this audience is not an edge case. It sits in the tap sheet ABOVE 'Log it' and outside the scroller, because the first version put it last in the scrolling body where it fell below the fold, and a route that has to be scrolled to is just a second hidden gesture."),

    ("A crash caught by reading rather than running", "Fixed", "2026-09-29",
     "A Reanimated worklet captures the variables it references AT THE MOMENT IT IS BUILT. clearCarry was declared below the gestures that used it, so building them would have read a const still in its temporal dead zone and thrown before the Week tab drew anything. Invisible to typecheck, lint and every screenshot of a screen at rest."),

    ("The hold was bound twice", "Fixed", "2026-09-29",
     "The pan gesture lifted the card AND Pressable's onLongPress threw a sheet up over the week she was dragging it across. Found in a mid-drag screenshot, which is the only place it could have been found. The hold belongs to the drag now; nothing is lost, because 'Move to…' is a tap."),

    ("Four pieces of polish, all hers", "Done", "2026-09-29",
     "The carried card was grey because it used backgroundSelected - the token for a selected tab - while every card on the screen uses backgroundElement, the brand sand. One token wrong and the thing under her finger looked like a different object from the thing she picked up. Today and the drop target were the same picture, both a filled row with an accentDeep border, so on a Tuesday you could not tell them apart: the drop is a new accentWash token and today is a dot. The section a card came from no longer lights on lift. Anytime never lights at all - the branch is gone, not merely false."),

    ("MY SCREENSHOT SCRIPT CHANGED HER DATA", "Fixed", "2026-09-29",
     "Clicking by fuzzy label match against her live account answered the Gym cadence question that was hers to answer, took Gym off Monday and Thursday, and moved Rocket yoga to Sunday. All three restored by hand and confirmed no logs were written. check-food-parse already carried the rule - a verification that changes the thing it is verifying is not a verification - and this script ignored it. Every mutating request to Supabase is now aborted at the network layer during a shoot."),

    ("Her own changes, correctly left alone", "Done", "2026-09-29",
     "Later the same rows changed again at 13:05 and 13:06 while she was testing on her phone. Every shoot since the guard reports writes blocked or none attempted, so those were hers - answering the cadence question and moving things about. Left exactly as she set them. Reverting her answers would have been worse than the original mistake."),

    ("check-week-drag.mjs, six checks", "Done", "2026-09-29",
     "Guards the worklet declaration order, the hold being bound once, activateAfterLongPress surviving, 'Move to…' staying in the tap sheet and outside the scroller, Anytime having no highlight branch, and isToday not choosing the day row's surface. Each proved to fail against a copy with its subject broken. The last two guard ABSENCES, which is what nobody notices coming back."),

    ("check-week-plan.mjs, sixteen checks", "Done", "2026-09-29",
     "placedOnDays across both guidance modes, the cadence conflict only when both answers are definite, the conservative log matcher and timesPerWeek declining to guess. Each group re-run against a stub that answers nothing, and the file reports a group as USELESS rather than green if the stub still passes it."),

    ("Merged to main", "Done", "2026-09-29",
     "After she confirmed drag on the phone, as she asked. All 17 check files, both typechecks, lint unchanged at 5 pre-existing errors with none in the new files."),

    ("Large text: STILL NOT TESTED", "Outstanding", "2026-09-29",
     "Unchanged from this morning and for the same reason. React Native declares sizes in points and the browser renders fixed pixels, so the web preview has no equivalent knob. Needs her phone."),

    ("Dropping into Anytime: verified by her, not by me", "Outstanding", "2026-09-29",
     "The automated drag test could not reach it - the Anytime block sits below the fold and the finger went off-screen. That it WORKS is known because she reported the highlight on it, which means she performed the drop. That it no longer highlights is guaranteed structurally rather than by screenshot: the branch does not exist."),
]

DECISIONS = [
    ("A tool that drives the real app against real data must be UNABLE to write", "2026-09-29",
     "The screenshot script signs in as Ruth and clicks by matching label text, and it changed her week. The obvious fix was a more careful selector, and that would have been the whole of it. But the rule was already written in this repository - check-food-parse stubs the client so a write is impossible - and the screenshot script had the same access with no such guard, because it felt like a camera rather than a client. Mutating requests are now aborted at the network layer. The printout of what was blocked turned out to be the useful half: a later run reported a blocked PATCH, and that line was the proof the drag had completed end to end, which no screenshot could show."),

    ("A substitution made quietly becomes a bug report she has to write", "2026-09-29",
     "She asked for long-press and drag. I built the hold as a shortcut to a chooser sheet, wrote the trade-off at the top of the file in four careful paragraphs, and then told her 'hold a card and move it to another day'. She read that as drag, because it is drag. The comment was not the problem - the comment was good. The trade-off was recorded where only I would read it and described to her in words that hid it. A decision to build something smaller than what was asked is hers to accept or refuse, and she can only do that if the message says so in the first sentence. The test: if she would be surprised to learn what the code actually does, the summary is wrong however true each of its words is."),

    ("Two states that mean different things must not share a surface", "2026-09-29",
     "Today and 'the card will land here' were both a filled backgroundSelected row with an accentDeep border, so on a Tuesday the drop target and today were indistinguishable. The first instinct was to differentiate by degree - a thicker border, a slightly deeper fill - which fails at a large font size and fails for anyone who cannot compare two things that are not on screen together. They differ in KIND now: the drop is the only filled row, and today is a dot. A dot and a tint cannot be confused at any size."),

    ("When a framework captures closures eagerly, declaration order is semantics", "2026-09-29",
     "A Reanimated worklet captures what it references at the moment it is BUILT, not when it runs, so a function declared below the gesture that uses it is read while still in its temporal dead zone. The screen throws before it draws. Nothing about this is visible to a typecheck or a lint, and no screenshot of a working screen would ever show it - it is only visible by reading, and only if you know the rule. Worth stating plainly because the instinct that 'order does not matter, these are all hoisted consts in one function' is correct everywhere else in the file."),
]


def main() -> int:
    wb = openpyxl.load_workbook(BOOK)
    check, dec = wb["Checklist"], wb["Decisions"]
    done_fill = copy(check.cell(479, 2).fill)
    header_font = copy(check.cell(477, 1).font)

    def append(ws, header, rows, status_col):
        r = ws.max_row + 2
        ws.cell(r, 1, header).font = copy(header_font)
        r += 1
        for row in rows:
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                if status_col and c == status_col and value.startswith(("Done", "Fixed", "Diagnosed")):
                    cell.fill = copy(done_fill)
            r += 1
        return r - 1

    check_last = append(check, LABEL, CHECKLIST, 2)
    dec_last = append(dec, LABEL, DECISIONS, None)
    wb.save(BOOK)

    back = openpyxl.load_workbook(BOOK)
    print(f"  sheets: {back.sheetnames}")
    print(f"  Checklist: {back['Checklist'].max_row} rows, label at {check_last - len(CHECKLIST)}")
    print(f"  Decisions: {back['Decisions'].max_row} rows, label at {dec_last - len(DECISIONS)}")
    print(f"  label: {LABEL}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
