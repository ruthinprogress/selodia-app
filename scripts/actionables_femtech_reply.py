"""Record the Accelerating FemTech reply on the funding table, in place.

Ruth, 2 October 2026, from her phone: "Just a quick update to log from
Accelerating FemTech."

WHAT CAME IN. HIN South London replied at 17:43 on 2 October to the enquiry she
sent on 29 September: "I have shared your message with the team and one of our
colleagues will be in touch with you soon." An acknowledgement from the
communications team, not an answer, and not from a programme lead.

WHY IT GOES ON THE EXISTING ROW RATHER THAN A NEW SECTION. The funding table
already carries "Accelerating FemTech: the email is drafted and NOT sent", which
is now two states out of date - it went on 29 September and has been answered. A
second row saying the same thing differently is how a list stops being readable.

IT APPENDS A SENTENCE, IT DOES NOT REWRITE THE CELL. Her words in that row are
hers, including the warning about which address can send. The reply is added to
the end of the note so nothing she wrote is lost.

AND IT EDITS NOTHING ELSE. Same discipline as actionables_funding.py: a dated
backup first, then one cell changed, found by its text rather than by its index.
"""

import shutil
import sys
from datetime import date
from pathlib import Path

import docx

DRIVE = Path(r"H:\My Drive\Selodia App Project Master Folder\Build Specs")
DOC = DRIVE / "Selodia - Open Actionables.docx"
BACKUP = DRIVE / f"Selodia - Open Actionables (backup {date.today():%Y-%m-%d} before femtech reply).docx"

# The row to update, matched on its opening words so a later edit to the rest of
# the sentence does not stop this finding it.
ROW_STARTS = "Accelerating FemTech"

ADDITION = (
    " SENT 29 SEPTEMBER, AND ANSWERED. HIN South London replied at 17:43 on "
    "2 October 2026: \u201cThank you for your email. I have shared your message with "
    "the team and one of our colleagues will be in touch with you soon.\u201d That is "
    "the communications team acknowledging it, not a programme lead answering it, "
    "so the thing still outstanding is a real reply about when the early-stage "
    "programme reopens. Worth a polite chase if nothing arrives by roughly "
    "16 October. The send address worked, which also settles the warning above."
)


def main() -> int:
    if not DOC.exists():
        print(f"  NOT FOUND: {DOC}")
        return 1

    d = docx.Document(str(DOC))

    # Find the cell by its text, across every table, rather than assuming which
    # table the funding section is.
    target = None
    for table in d.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip().startswith(ROW_STARTS):
                target = row
                break
        if target is not None:
            break

    if target is None:
        print(f"  NOT FOUND: no row starting '{ROW_STARTS}'. Nothing written.")
        return 1

    note = target.cells[-1]
    if "ANSWERED" in note.text:
        print("  already recorded. Nothing written.")
        return 0

    shutil.copy2(DOC, BACKUP)
    print(f"  backed up to: {BACKUP.name}")

    # Appended to the LAST run of the last paragraph, so the cell keeps its own
    # formatting instead of gaining a differently-styled block.
    para = note.paragraphs[-1]
    if para.runs:
        para.runs[-1].text = para.runs[-1].text + ADDITION
    else:
        para.add_run(ADDITION)

    d.save(str(DOC))
    print("  recorded on the funding table:")
    print(f"    {target.cells[0].text.strip()[:70]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
