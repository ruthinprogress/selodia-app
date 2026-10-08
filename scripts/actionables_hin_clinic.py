"""Close the HIN enquiry row and open the innovation clinic row.

Ruth forwarded the reply on 7 October 2026. Jade Okparanta, Programme Support
Officer at Health Innovation Network South London, answered the enquiry she sent
on 29 September.

WHAT IT ACTUALLY SAYS, which is three things and only one of them is an opening:

  There is no date for the next programme and no list to join. That closes the
  row, because the thing it was waiting for has arrived and the answer is no.

  Newsletters are the only way to hear about the next one.

  There is a free innovation clinic, booked through an Innovator support form.
  That is a new row, not a continuation of this one: a different thing, with a
  different ask, that she has to do.

IT APPENDS AND IT DOES NOT REWRITE. Her words in that row are hers. The reply
goes on the end of the note so nothing she wrote is lost, which is the same
discipline as actionables_femtech_reply.py.

NO BACKUP COPY, AND THAT IS A CHANGE FROM THE EARLIER SCRIPTS. docs/where-things-go.md:
"Never write a backup copy next to a file you are editing. H:\\My Drive is Google
Drive, and Drive keeps version history on every file in it." The older scripts in
this folder predate that rule and should not be copied on this point.
"""

import sys
from copy import deepcopy
from pathlib import Path

import docx

DOC = Path(
    r"H:\My Drive\Selodia App Project Master Folder\Build Specs"
) / "Selodia - Open Actionables.docx"

# Matched on opening words so a later edit to the rest of the sentence does not
# stop this finding it.
FEMTECH_STARTS = "Accelerating FemTech"

CLOSED = (
    " ANSWERED AND CLOSED, 7 October 2026. Jade Okparanta, Programme Support "
    "Officer, replied: there is no set date for the next programme and no list to "
    "sign up to for it. The newsletters for DigitalHealth.London and HIN South "
    "London are the only way to hear about the next one, so signing up to both is "
    "the whole of what can be done here. Nothing further is outstanding on this "
    "row. The same reply offered a free innovation clinic, which is its own row "
    "below because it is a different thing with a different ask."
)

CLINIC_WHAT = "HIN innovation clinic: the form is Ruth's to send"

CLINIC_NOTE = (
    "Offered unprompted in Jade Okparanta's reply of 7 October 2026: a free "
    "clinic giving advice and signposting on the innovation, booked by completing "
    "their Innovator support form. Ruth has never registered with the NHS "
    "Innovation Service and has no engagement form with any Health Innovation "
    "Network, so there is no Innovation Record Number to quote and it is the full "
    "form. The content is drafted and waiting in Build Specs / 2026-10-07 HIN "
    "innovation clinic - what to send and what to ask.docx. Code did not submit "
    "it: it asks for company details and goes to an outside body. "
    "THE QUESTION WORTH THE WHOLE CLINIC is whether the eleven symptom detectors "
    "put any part of Selodia in scope as Software as a Medical Device. Six route "
    "to 999 and five to 111, they are approved by Ruth and by no clinician, and a "
    "product that takes a description of symptoms and points somebody toward "
    "emergency care is in the territory where that question gets asked, whatever "
    "the answer turns out to be. HIN South London and DigitalHealth.London have "
    "run a programme on exactly that regulatory route, so they are the right "
    "people to ask and it costs nothing. "
    "THE SECOND ASK is a route to a clinical reviewer for those detectors, which "
    "is already on this list at GBP 600 to 1,200, gates the beta, and would also "
    "satisfy the 14 November evidence-of-a-team row. "
    "WORTH DOING THIS WEEK rather than in November: these take time to book, and "
    "the answer is most useful before the Women in Innovation round opens, which "
    "is expected around 26 November on a pattern of one year."
)


def styled_cell(cell, template_cell, text):
    """Write `text` into `cell`, wearing the formatting of `template_cell`."""
    source = template_cell.paragraphs[0]
    target = cell.paragraphs[0]
    # Copy the paragraph's own formatting, then give it one run that looks like
    # the template's first run. Building a cell from scratch produces a row in
    # the document's default face, which is visibly not the table's.
    target._p.getparent().replace(target._p, deepcopy(source._p))
    para = cell.paragraphs[0]
    for run in para.runs[1:]:
        run._r.getparent().remove(run._r)
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def main() -> int:
    if not DOC.exists():
        print(f"  NOT FOUND: {DOC}")
        return 1

    doc = docx.Document(DOC)

    target = None
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            if row.cells[0].text.strip().startswith(FEMTECH_STARTS):
                target = (table, i, row)
                break
        if target:
            break

    if not target:
        print(f'  NOT FOUND: no row starting "{FEMTECH_STARTS}". Nothing written.')
        return 1

    table, index, row = target

    note = row.cells[1]
    if "ANSWERED AND CLOSED" in note.text:
        print("  Already closed. Nothing written.")
        return 0

    para = note.paragraphs[-1]
    if para.runs:
        para.runs[-1].text = para.runs[-1].text + CLOSED
    else:
        para.add_run(CLOSED)
    print(f"  closed: {row.cells[0].text.strip()[:60]}")

    if any(r.cells[0].text.strip().startswith("HIN innovation clinic") for r in table.rows):
        print("  The clinic row already exists. Not added twice.")
    else:
        # add_row appends; addnext moves it to sit directly under the row it
        # follows from, which is where somebody reading the list would look.
        new = table.add_row()
        row._tr.addnext(new._tr)
        styled_cell(new.cells[0], row.cells[0], CLINIC_WHAT)
        styled_cell(new.cells[1], row.cells[1], CLINIC_NOTE)
        print(f"  added:  {CLINIC_WHAT}")

    doc.save(DOC)

    # READ BACK, because a save that silently wrote nothing is the failure this
    # kind of script produces most often.
    check = docx.Document(DOC)
    closed = any(
        "ANSWERED AND CLOSED" in r.cells[1].text
        for t in check.tables
        for r in t.rows
        if r.cells[0].text.strip().startswith(FEMTECH_STARTS)
    )
    opened = any(
        r.cells[0].text.strip().startswith("HIN innovation clinic")
        for t in check.tables
        for r in t.rows
    )
    print(f"  read back: closed={closed} clinic_row={opened}")
    return 0 if (closed and opened) else 1


if __name__ == "__main__":
    sys.exit(main())
