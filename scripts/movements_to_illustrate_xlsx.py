"""Write the "movements to illustrate" list into Ruth's Drive as a spreadsheet.

    python scripts/movements_to_illustrate_xlsx.py

Ruth, 18 October 2026: "Please turn that list into a spreadsheet that
autopopulates, as I doubt Ill ever look at it otherwise."

WHICH IS THE WHOLE POINT. what-to-make-next.mjs prints the same list in a
terminal, and a terminal is somewhere she does not go. A list nobody reads
cannot decide anything, so it goes where her other working documents are and
rewrites itself.

THE SPREADSHEET IS A VIEW, NOT THE RECORD. movement_gaps in the database is the
record; this file is regenerated from it every time and may be deleted and
remade without losing anything. That is why it is safe to overwrite in place
rather than making the dated copies that cluttered this folder in October.
"""

import io
import json
import os
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
OUT = Path(
    r"H:\My Drive\Selodia App Project Master Folder\Build Specs"
    r"\Selodia - Movements To Illustrate.xlsx"
)

INK, ACCENT, MUTED = "2D2B28", "874C3A", "605A52"
BAND = "F7F3EA"


def env():
    values = {}
    for line in (REPO / ".env.local").read_text(encoding="utf-8").splitlines():
        i = line.find("=")
        if i > 0 and not line.startswith("#"):
            values[line[:i].strip()] = line[i + 1 :].strip().strip("\"'")
    return values


def fetch(e):
    url = (
        e["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/")
        + "/rest/v1/movement_gaps"
        + "?select=name,times_asked,first_seen,last_seen,last_plan,status"
        + "&order=status.asc,times_asked.desc,last_seen.desc"
    )
    key = e["SUPABASE_SERVICE_ROLE_KEY"]
    req = urllib.request.Request(
        url, headers={"apikey": key, "Authorization": f"Bearer {key}"}
    )
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))


def day(stamp):
    if not stamp:
        return ""
    return str(stamp)[:10]


def main() -> int:
    e = env()
    rows = fetch(e)
    wanted = [r for r in rows if r.get("status") == "wanted"]
    done = [r for r in rows if r.get("status") != "wanted"]

    wb = openpyxl.Workbook()

    # ---- Sheet 1: the list itself ----
    ws = wb.active
    ws.title = "To illustrate"

    ws["A1"] = "Movements to illustrate"
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws["A2"] = (
        "Every movement somebody was given as written instructions because the app "
        "had no film and no drawing of it."
    )
    ws["A2"].font = Font(size=10, color=MUTED)
    ws["A3"] = (
        "Ordered by how many times it has been asked for. The top of the list is the "
        "thing most people keep missing."
    )
    ws["A3"].font = Font(size=10, color=MUTED)
    ws["A4"] = (
        f"Rewritten automatically. Last updated "
        f"{datetime.now(timezone.utc).strftime('%d %B %Y, %H:%M')} UTC. "
        f"See the 'How this works' sheet."
    )
    ws["A4"].font = Font(size=10, italic=True, color=ACCENT)

    headers = ["Times asked", "Movement", "First asked", "Last asked", "Most recently for"]
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=6, column=col, value=text)
        c.font = Font(bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=BAND)

    r = 7
    for row in wanted:
        ws.cell(row=r, column=1, value=row.get("times_asked"))
        ws.cell(row=r, column=2, value=row.get("name"))
        ws.cell(row=r, column=3, value=day(row.get("first_seen")))
        ws.cell(row=r, column=4, value=day(row.get("last_seen")))
        ws.cell(row=r, column=5, value=row.get("last_plan"))
        r += 1

    if not wanted:
        ws.cell(row=7, column=2, value="Nothing outstanding.").font = Font(italic=True, color=MUTED)
        r = 8

    if done:
        r += 1
        c = ws.cell(row=r, column=1, value="Already made")
        c.font = Font(bold=True, color=ACCENT)
        r += 1
        for row in done:
            ws.cell(row=r, column=1, value=row.get("times_asked"))
            ws.cell(row=r, column=2, value=row.get("name"))
            ws.cell(row=r, column=3, value=day(row.get("first_seen")))
            ws.cell(row=r, column=4, value=day(row.get("last_seen")))
            ws.cell(row=r, column=5, value=row.get("status"))
            r += 1

    for col, width in zip("ABCDE", (12, 46, 13, 13, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"

    # ---- Sheet 2: how it got here ----
    how = wb.create_sheet("How this works")
    lines = [
        ("Movements to illustrate - how this works", "title"),
        ("", None),
        ("WHY IT EXISTS", "head"),
        (
            "The app teaches a movement in one of three ways: a filmed clip, a drawing, or "
            "written instructions. All three are real and a plan containing any of them is "
            "complete. But a movement that only ever arrives as words is one we could teach "
            "better, and this is the list of those - in the order people keep asking for them.",
            None,
        ),
        (
            "It replaces guessing. The library grows toward what people actually came for "
            "rather than toward whoever last had an idea.",
            None,
        ),
        ("", None),
        ("HOW IT IS BUILT", "head"),
        (
            "Nobody adds to this by hand. When a movement plan is saved, the app looks up "
            "every movement in it to find its clip. Anything that comes back with neither a "
            "clip nor a drawing is written down at that moment, and the count goes up by one.",
            None,
        ),
        (
            "The chat is never asked to flag a gap, so it cannot forget to. The recording "
            "happens in the code that writes the plan - app/lib/movement-gaps.ts.",
            None,
        ),
        (
            "It counts across EVERY account, not just Ruth's. There is no user column on the "
            "table: one row per movement, shared by everyone. As the app grows this becomes "
            "real demand from real users.",
            None,
        ),
        ("", None),
        ("HOW A ROW LEAVES THE LIST", "head"),
        (
            "When the illustration is uploaded, the row is marked as made and drops to the "
            "'Already made' section at the bottom. That happens inside the upload script, so "
            "it cannot be forgotten either.",
            None,
        ),
        (
            "The picture then appears by itself in plans people already have. Movement cards "
            "look up what exists every time they open, rather than remembering what existed "
            "when the plan was written - so somebody who asked for a movement sees it turn up "
            "in their own plan once it is made.",
            None,
        ),
        ("", None),
        ("WHERE EVERYTHING LIVES", "head"),
        ("This file:  Build Specs \\ Selodia - Movements To Illustrate.xlsx", None),
        ("The record: the movement_gaps table in Supabase. This file is only a view of it,", None),
        ("            rewritten each time, and safe to delete.", None),
        ("Rebuild it: python scripts/movements_to_illustrate_xlsx.py", None),
        ("Same list in a terminal: node scripts/what-to-make-next.mjs", None),
        ("", None),
        ("The drawings themselves: Build Specs \\ Branding & Assets \\ Visual Assets", None),
        ("How to draw one: scripts/illustration-brief.mjs holds the art direction and every", None),
        ("            rule learned the expensive way. Prompts are built by movementPrompt().", None),
    ]
    r = 1
    for text, kind in lines:
        c = how.cell(row=r, column=1, value=text)
        if kind == "title":
            c.font = Font(bold=True, size=16, color=INK)
        elif kind == "head":
            c.font = Font(bold=True, color=ACCENT)
        else:
            c.font = Font(size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    how.column_dimensions["A"].width = 104

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"\n  {OUT}")
    print(f"  {len(wanted)} to illustrate, {len(done)} already made\n")
    for row in wanted[:5]:
        print(f"    {row['times_asked']:>3}  {row['name']}")
    if len(wanted) > 5:
        print(f"    ... and {len(wanted) - 5} more\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
